# -*- coding: utf-8 -*-
"""verify_rollup.py - независимый пересчёт (Python, без Excel) итогов ГРЩ по категориям из книг объекта.

    py verify_rollup.py "<папка объекта>" [--grsh "Однолинейка ГРЩ-ЖД.xlsx"]

Читает кэш значений (openpyxl, data_only): линии каждого щита («Нагрузка_щитов»: группа A, мощность H, Кс сил B;
«Исходные данные»: категория BP, пожарная BQ), таблицу Кс («Справочная» BP4:BQ11) и Ко (BQ14). Собирает объект так же, как лист «Режимы»
(см. 4_Инструкции/Категории_режимы_ИБП.md, раздел 6), но по сырым линиям, а не по формулам сборки:

  собственные линии ГРЩ + линии нижестоящих щитов; на каждой линии-фидере нижнего щита его нагрузка раскладывается по
  категориям и пожарности его линий; если линия-фидер помечена пожарной, вся нагрузка нижнего щита идёт в пожарные
  (в категорию линии-фидера). Рр набора = Руст быт·Кс(табл.) + Ко·Σ(Pсил·Кс) + Σ Pпост.

Сравнивает с листом «Режимы» ГРЩ (D103:I107). Код возврата 0 - совпало (допуск 1e-6), 1 - расхождение.
"""
import argparse
import glob
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
CAT = ["I особая", "I", "II", "III"]
TOL = 1e-6


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


class Book:
    def __init__(self, path):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=True)

    def lines(self):
        """линии щита, попавшие в расчёт: «Нагрузка_щитов» D (номер линии), A (группа), H (мощность), B (Кс сил),
        категория и пожарная - «Исходные данные» BP, BQ"""
        sd = self.wb["Исходные данные"]
        cat = {}
        for r in range(3, 1001):
            d = sd["D%d" % r].value
            if d not in (None, ""):
                cat[str(d)] = (sd["BP%d" % r].value or "III", sd["BQ%d" % r].value == "да")
        ws = self.wb["Нагрузка_щитов"]
        out = []
        for r in range(3, 1001):
            d, g = ws["D%d" % r].value, ws["A%d" % r].value
            if d in (None, "") or g in (None, ""):
                continue
            ks = ws["B%d" % r].value
            out.append(dict(name=str(d), e=num(ws["H%d" % r].value) or 0.0, grp=g, ks=num(ks) if g == "Сил" else None,
                            cat=cat.get(str(d), ("III", False))[0], fire=cat.get(str(d), ("III", False))[1]))
        return out

    def ks_table(self):
        ws = self.wb["Справочная"]
        xs = [num(ws["BP%d" % r].value) for r in range(4, 12)]
        ys = [num(ws["BQ%d" % r].value) for r in range(4, 12)]
        return xs, ys, num(ws["BQ14"].value)


def ks_byt(x, xs, ys):
    if x <= 0:
        return 0.0
    if x >= xs[-1]:
        return ys[-1]
    i = max(k for k in range(len(xs)) if xs[k] <= x)
    return ys[i] + (x - xs[i]) * (ys[i + 1] - ys[i]) / (xs[i + 1] - xs[i])


def comps(line):
    """(Руст, быт, Σ P·Кс сил, пост) одной линии"""
    g, e = line["grp"], line["e"]
    if g in ("Нет", "", "Щит"):
        return (0.0, 0.0, 0.0, 0.0) if g != "Щит" else (e, 0.0, 0.0, 0.0)
    return (e, e if g == "Быт" else 0.0, e * (line["ks"] if line["ks"] is not None else 1.0) if g == "Сил" else 0.0,
            e if g == "Пост" else 0.0)


def rr(c, xs, ys, ko):
    return c[1] * ks_byt(c[1], xs, ys) + ko * c[2] + c[3]


def add(a, b):
    return tuple(x + y for x, y in zip(a, b))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--grsh", default="Однолинейка ГРЩ-ЖД.xlsx")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    top = Book(os.path.join(a.folder, a.grsh))
    xs, ys, ko = top.ks_table()
    nh = top.wb["Нижестоящие_щиты"]
    feeders = {}
    for r in range(5, 35):
        b = nh["B%d" % r].value
        f = nh["C%d" % r].value
        if b and f and os.path.exists(os.path.join(a.folder, f)):
            feeders[str(b)] = f
    z = (0.0,) * 4
    nf = {j: z for j in CAT}         # без пожарных: по категории
    fr = {j: z for j in CAT}         # пожарные: по категории линии-фидера / своей линии
    own_n = 0
    lower_n = 0
    shield_rows = []
    for ln in top.lines():
        if ln["grp"] == "Щит":
            fdr = feeders.get(ln["name"])
            if not fdr:
                print("линия-фидер %s без файла нижестоящего щита: пропущена" % ln["name"])
                continue
            bk = Book(os.path.join(a.folder, fdr))
            row = dict(name=ln["name"], nf={j: z for j in CAT}, fire=z)
            for sl in bk.lines():
                c = comps(sl)
                if ln["fire"] or sl["fire"]:
                    fr[ln["cat"]] = add(fr[ln["cat"]], c)
                    row["fire"] = add(row["fire"], c)
                else:
                    nf[sl["cat"]] = add(nf[sl["cat"]], c)
                    row["nf"][sl["cat"]] = add(row["nf"][sl["cat"]], c)
            shield_rows.append(row)
            lower_n += 1
        else:
            c = comps(ln)
            if ln["fire"]:
                fr[ln["cat"]] = add(fr[ln["cat"]], c)
            else:
                nf[ln["cat"]] = add(nf[ln["cat"]], c)
            own_n += 1
    print("ГРЩ: собственных линий %d, нижестоящих щитов %d" % (own_n, lower_n))
    # ожидания
    exp = {}
    for j in CAT:
        exp[j] = (nf[j][0], rr(nf[j], xs, ys, ko), fr[j][0], rr(fr[j], xs, ys, ko))
    tot_n = tuple(sum(nf[j][k] for j in CAT) for k in range(4))
    tot_f = tuple(sum(fr[j][k] for j in CAT) for k in range(4))
    exp["Итого"] = (tot_n[0], rr(tot_n, xs, ys, ko), tot_f[0], rr(tot_f, xs, ys, ko))
    # ГРЩ, «Режимы»
    m = top.wb["Режимы"]
    rows = {"I особая": 103, "I": 104, "II": 105, "III": 106, "Итого": 107}
    ok = True
    print("%-10s %-12s %-12s %-12s %-12s | %s" % ("", "Руст", "Рр", "пож.Руст", "пож.Рр", "разница с «Режимами» (Руст, Рр, пож.Руст, пож.Рр)"))
    for k, r in rows.items():
        got = (num(m["D%d" % r].value) or 0.0, num(m["E%d" % r].value) or 0.0, num(m["H%d" % r].value) or 0.0, num(m["I%d" % r].value) or 0.0)
        e = exp[k]
        d = [g - x for g, x in zip(got, e)]
        bad = any(abs(x) > TOL for x in d)
        ok &= not bad
        print("%-10s %-12.4f %-12.4f %-12.4f %-12.4f | %s %s" % (k, e[0], e[1], e[2], e[3], " ".join("%9.6f" % x for x in d), "РАСХОЖДЕНИЕ" if bad else "ок"))
    # по щитам (Руст, лист «Категории_по_щитам»)
    s = top.wb["Категории_по_щитам"] if "Категории_по_щитам" in top.wb.sheetnames else None
    if s is not None:
        print("по щитам (Руст по категориям и пожарные; лист «Категории_по_щитам» C:G):")
        got = {}
        for r in range(19, 49):
            nm = s["B%d" % r].value
            if nm:
                got[str(nm)] = [num(s["%s%d" % (c, r)].value) or 0.0 for c in "CDEFG"]
        for row in shield_rows:
            e = [row["nf"][j][0] for j in CAT] + [row["fire"][0]]
            g = got.get(row["name"])
            if g is None:
                print("  %-12s нет строки" % row["name"])
                ok = False
                continue
            # пожарные фидера в ячейке «Пожарные» - Руст пожарных нижнего щита
            bad = any(abs(x - y) > TOL for x, y in zip(g, e))
            ok &= not bad
            if bad or any(v for v in e[:3]) or e[4]:
                print("  %-12s ожидание %s | лист %s %s" % (row["name"], [round(x, 3) for x in e], [round(x, 3) for x in g], "РАСХОЖДЕНИЕ" if bad else "ок"))
    print("ИТОГ:", "совпало" if ok else "РАСХОЖДЕНИЕ")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
