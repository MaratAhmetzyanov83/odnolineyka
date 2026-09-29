"""Перевод однолинеек объекта «Однолинейки ЖД» (старый формат, 11 листов) в формат v3.11 + связи ГРЩ ↔ нижестоящие щиты.

python3 convert_zhd.py <папка исходников> <папка результата>
Шаги: 1) каждый дочерний файл → шаблон v3.11 (данные «Исходных данных», «Разбивки», ручные фазы, коды и этажи объекта,
группы Кс и категории ПЗ); 2) пересчёт LibreOffice → экспорт каждого щита; 3) ГРЩ: внешние ссылки на экспорт щитов
(«Нижестоящие_щиты» D:J) с кэшем значений; 4) пересчёт ГРЩ → K:N; 5) щиты: внешние ссылки на K:N ГРЩ («Питание_щита» C6:F6).
"""
import os, re, sys, shutil, subprocess, html, urllib.parse, json
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import patch7 as P7
from patch7 import X, rd, wr, ct, cn, cf, ce, bulk, unshare, esc
import openpyxl

TPL = os.path.join(HERE, "out11", "Однолинейка_шаблон.xlsx")
GRSH = os.path.join(HERE, "out11", "Однолинейка ГРЩ-ЖД.xlsx")
SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
WORK = os.path.join(HERE, "zhd_work"); os.makedirs(WORK, exist_ok=True)
OBJ = "Резиденция, здание жилого дома"

# ГРЩ: линия → файл щита, Кс силовой нагрузки из таблицы ГРЩ (его прежнее допущение)
LINKS = [("ЩОВ-01", "Однолинейка_ЩОВ-01.xlsx", 0.5), ("ЩОВ-001", "Однолинейка_ЩОВ-001.xlsx", 0.5),
         ("ЩОВ-002", "Однолинейка_ЩОВ-002.xlsx", 0.5), ("ЩОВ-2", "Однолинейка_ЩОВ-2.xlsx", 0.5),
         ("ЩР-001", "Однолинейка_ЩР-001.xlsx", None), ("ЩР-01", "Однолинейка_ЩР-01.xlsx", None),
         ("ЩР-1", "Однолинейка_ЩР-1.xlsx", None), ("ЩР-2", "Однолинейка_ЩР-2.xlsx", None),
         ("ЩС-Б", "Однолинейка_ЩС-Б.xlsx", 0.5), ("ЩС-ВК.001", "Однолинейка_ЩС-ВК.001.xlsx", 0.5),
         ("ЩС-ВК.002", "Однолинейка_ЩС-ВК.002.xlsx", 0.7), ("ЩС-ВК.003", "Однолинейка_ЩС-ВК.003.xlsx", 0.7),
         ("ЩУОВ", "Однолинейка_ЩУОВ.xlsx", None), ("ЩСПЗ-ЖД.1", "Однолинейка_ЩСПЗ-ЖД.xlsx", None)]
KS_BY_FILE = {f: k for _, f, k in LINKS}
# группа по умолчанию для всего щита (кроме вводов)
PANEL_GROUP = {"ЩОВ": "Сил", "ЩС-Б": "Сил", "ЩС-ВК": "Сил", "ЩУОВ": "Пост", "ЩСПЗ": "Пост"}
CAT = {"В": "Вентиляция", "П": "Вентиляция", "ПВ": "Вентиляция", "V": "Вентиляция",
       "K": "Кондиционирование и увлажнение", "К": "Кондиционирование и увлажнение", "ККБ": "Кондиционирование и увлажнение",
       "НБ": "Кондиционирование и увлажнение", "У": "Кондиционирование и увлажнение",
       "R": "Розеточная сеть и бытовые приборы", "CR": "Розеточная сеть и бытовые приборы", "FR": "Розеточная сеть и бытовые приборы",
       "L": "Рабочее освещение", "LD": "Рабочее освещение", "LT": "Рабочее освещение", "DA": "Рабочее освещение",
       "FL": "Наружное и фасадное освещение", "FLD": "Наружное и фасадное освещение",
       "ШТК": "Слаботочные системы, автоматизация, KNX", "KNX": "Слаботочные системы, автоматизация, KNX",
       "ЩКД": "Системы безопасности и связи", "ЩДК": "Системы безопасности и связи", "ШКД": "Системы безопасности и связи",
       "Э": "Бассейн, SPA, хаммам", "ЩС": "Бассейн, SPA, хаммам", "НВК": "Насосы (водоснабжение, дренаж)",
       "ЖУ": "Насосы (водоснабжение, дренаж)", "Н": "Насосы (водоснабжение, дренаж)", "СПД": "Насосы (водоснабжение, дренаж)",
       "OV": "Обогрев водостоков и кровли", "DT": "Обогрев водостоков и кровли", "DW": "Обогрев водостоков и кровли",
       "ЭОК": "Обогрев водостоков и кровли", "FCP": "Системы противопожарной защиты", "ENC": "Системы противопожарной защиты",
       "ЩО": "Рабочее освещение"}
POST_PREFIX = ("ШТК", "ЩКД", "ЩДК", "ШКД", "KNX", "EIB", "WD", "VK", "ТЕ", "РМ")
UGO = {"R": ("Розетка", "да", 16), "CR": ("Розетка", "да", 16), "FR": ("Розетка", "да", 16),
       "L": ("Светильник", "нет", 10), "LD": ("Светильник", "нет", 10), "LT": ("Светильник", "нет", 10),
       "FL": ("Светильник", "нет", 10), "FLD": ("Светильник", "нет", 10), "DA": ("Светильник", "нет", 10),
       "SH": ("Привод (двигатель)", "нет", None), "V": ("Привод (двигатель)", "нет", None), "П": ("Привод (двигатель)", "нет", None),
       "В": ("Привод (двигатель)", "нет", None), "ПВ": ("Привод (двигатель)", "нет", None), "НБ": ("Привод (двигатель)", "нет", None),
       "К": ("Привод (двигатель)", "нет", None), "У": ("Привод (двигатель)", "нет", None), "Н": ("Привод (двигатель)", "нет", None),
       "СПД": ("Привод (двигатель)", "нет", None), "КЛ": ("Клапан", "нет", None), "ЭК": ("Клапан", "нет", None),
       "ЭОК": ("Нагреватель", "нет", None)}


def prefix(line):
    s = str(line).split(".")[0]
    m = re.match(r"^([A-Za-zА-Яа-яЁё]+)", s)
    return m.group(1) if m else s


def group_for(fn, line):
    p = prefix(line)
    if p in ("Ввод",) or str(line).startswith(("На ИБП", "От ИБП")): return "Нет"
    for k, g in PANEL_GROUP.items():
        if os.path.basename(fn).startswith("Однолинейка_" + k): return g
    if p.startswith(POST_PREFIX): return "Пост"
    return "Быт"


def val_xml(ref, v, st):
    if v is None or v == "": return ce(ref, st)
    if hasattr(v, "text"): v = v.text if str(v.text).startswith("=") else "=" + str(v.text)
    if isinstance(v, str) and v.startswith("="): return cf(ref, v[1:], st)
    if isinstance(v, bool): return cn(ref, int(v), st)
    if isinstance(v, (int, float)): return cn(ref, v, st)
    return ct(ref, v, st)


def convert_child(fn, out):
    src = openpyxl.load_workbook(os.path.join(SRC, fn))
    si, sr, sl, ss = src["Исходные данные"], src["Разбивка_по_щитам"], src["Нагрузка_щитов"], src["Справочная"]
    WD = os.path.join(WORK, "w_" + re.sub(r"\W", "_", fn)); order = P7.unpack(TPL, WD)
    X.load_sst(WD)
    # --- Справочная: коды объекта, этажи, марки кабеля
    spr = X.Sheet(X.sheet_path(WD, "Справочная")); c = {}
    for r in range(4, 52):
        code = ss.cell(r, 16).value
        for j, col in enumerate("PQRST"):
            c[f"{col}{r}"] = val_xml(f"{col}{r}", ss.cell(r, 16 + j).value, spr.style(f"{col}{r}"))
        for col in "UVWXY":
            c[f"{col}{r}"] = ce(f"{col}{r}", spr.style(f"{col}{r}"))
        u = UGO.get(str(code).strip()) if code else None
        c[f"AV{r}"] = val_xml(f"AV{r}", u[0] if u else None, spr.style(f"AV{r}"))
        c[f"AW{r}"] = val_xml(f"AW{r}", u[1] if u else None, spr.style(f"AW{r}"))
        c[f"AX{r}"] = val_xml(f"AX{r}", u[2] if u else None, spr.style(f"AX{r}"))
        c[f"AY{r}"] = val_xml(f"AY{r}", 2 if code == "SH" else None, spr.style(f"AY{r}"))
        c[f"Z{r}"] = val_xml(f"Z{r}", CAT.get(prefix(code)) if code else None, spr.style(f"Z{r}"))
        c[f"AZ{r}"] = ce(f"AZ{r}", spr.style(f"AZ{r}"))
    for r in range(4, 20):
        for j, col in enumerate(["AD", "AE", "AF", "AG"]):
            c[f"{col}{r}"] = val_xml(f"{col}{r}", ss.cell(r, 25 + j).value, spr.style(f"{col}{r}"))
    have = {X.value(spr, f"L{r}") for r in range(4, 52)}
    free = [r for r in range(4, 52) if not X.value(spr, f"L{r}")]
    for r in range(4, 52):
        code = ss.cell(r, 12).value
        if code and code not in have:
            rr = free.pop(0); have.add(code)
            c[f"L{rr}"] = ct(f"L{rr}", code, spr.style("L16")); c[f"M{rr}"] = ct(f"M{rr}", ss.cell(r, 13).value or "", spr.style("M16"))
            if ss.cell(r, 14).value is not None: c[f"N{rr}"] = cn(f"N{rr}", ss.cell(r, 14).value, spr.style("N16"))
    bulk(spr, c); spr.save()
    # --- Исходные данные: входные столбцы как в исходнике (строки 3–1000)
    isx = X.Sheet(X.sheet_path(WD, "Исходные данные"))
    COLS = ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "N", "O", "Q", "R", "S"]
    unshare(isx, set(COLS))
    c = {}; lines = []
    ks = KS_BY_FILE.get(fn)
    for r in range(3, 1001):
        for col in COLS:
            c[f"{col}{r}"] = val_xml(f"{col}{r}", si[f"{col}{r}"].value, isx.style(f"{col}{r}"))
        for col in ("AU", "AV", "AW", "AX", "AY"):
            if not isx.is_empty(f"{col}{r}"): c[f"{col}{r}"] = ce(f"{col}{r}", isx.style(f"{col}{r}"))
        ln = si[f"D{r}"].value
        if ln:
            g = group_for(fn, ln); lines.append((ln, g))
            c[f"AZ{r}"] = ct(f"AZ{r}", g, isx.style(f"AZ{r}"))
            if g == "Сил" and ks: c[f"BA{r}"] = cn(f"BA{r}", ks, isx.style(f"BA{r}"))
            cat = CAT.get(prefix(ln))
            if cat: c[f"BD{r}"] = ct(f"BD{r}", cat, isx.style(f"BD{r}"))
    bulk(isx, c); isx.save()
    # --- Разбивка, ручные фазы, «авто» в однолинейке
    rz = X.Sheet(X.sheet_path(WD, "Разбивка_по_щитам")); c = {}
    rows_with_line = []
    for r in range(3, 1001):
        for col in "FGH":
            v = sr[f"{col}{r}"].value
            c[f"{col}{r}"] = val_xml(f"{col}{r}", v, rz.style(f"{col}{r}"))
        if sr[f"H{r}"].value: rows_with_line.append(r)
    bulk(rz, c); rz.save()
    ld = X.Sheet(X.sheet_path(WD, "Нагрузка_щитов")); c = {}
    for r in range(3, 1001):
        v = sl[f"F{r}"].value
        c[f"E{r}"] = val_xml(f"E{r}", v, ld.style(f"E{r}"))
    bulk(ld, c); ld.save()
    ol = X.Sheet(X.sheet_path(WD, "Однолинейка")); c = {}
    for r in rows_with_line:
        c[f"E{r}"] = ct(f"E{r}", "авто", ol.style(f"E{r}"))
    bulk(ol, c); ol.save()
    panel = sorted({sr[f"F{r}"].value for r in range(3, 1001) if sr[f"F{r}"].value})
    gd = X.Sheet(X.sheet_path(WD, "Краткое руководство"))
    bulk(gd, {"C1": ct("C1", f"{', '.join(panel)} — однолинейка v3.11. Перенесено из «{fn}» (старый формат) 29.09.2026.", gd.style("C1") or gd.style("C2"))}); gd.save()
    pz = X.Sheet(X.sheet_path(WD, "Расчёт_нагрузок"))
    bulk(pz, {"C3": ct("C3", OBJ, pz.style("C3"))}); pz.save()
    P7.pack(WD, order, out)
    return WD, order, panel, lines


# ------------------------------------------------------------------ внешние ссылки
def add_ext_links(WD, order, links):
    """links: [(имя файла, лист, {ячейка: значение})] → номера [1..n] в том же порядке"""
    wbp = WD + "/xl/workbook.xml"; wb = rd(wbp)
    rp = WD + "/xl/_rels/workbook.xml.rels"; rels = rd(rp)
    ctp = WD + "/[Content_Types].xml"; ctx = rd(ctp)
    assert "<externalReferences>" not in wb
    os.makedirs(WD + "/xl/externalLinks/_rels", exist_ok=True)
    refs = []
    nid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', rels))
    for k, (fname, sheet, cache) in enumerate(links, 1):
        nid += 1; rid = f"rId{nid}"
        rows = {}
        for ref, v in cache.items():
            r = int(re.search(r"\d+", ref).group(0)); rows.setdefault(r, []).append((ref, v))
        sd = ""
        for r in sorted(rows):
            cells = ""
            for ref, v in sorted(rows[r], key=lambda x: X.col2n(re.match(r"[A-Z]+", x[0]).group(0))):
                if isinstance(v, (int, float)) and not isinstance(v, bool): cells += f'<cell r="{ref}"><v>{v}</v></cell>'
                else: cells += f'<cell r="{ref}" t="str"><v>{esc("" if v is None else v)}</v></cell>'
            sd += f'<row r="{r}">{cells}</row>'
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
               '<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1">'
               f'<sheetNames><sheetName val="{esc(sheet)}"/></sheetNames>'
               f'<sheetDataSet><sheetData sheetId="0">{sd}</sheetData></sheetDataSet></externalBook></externalLink>')
        part = f"xl/externalLinks/externalLink{k}.xml"; relp = f"xl/externalLinks/_rels/externalLink{k}.xml.rels"
        wr(os.path.join(WD, part), xml)
        tgt = urllib.parse.quote(fname)
        wr(os.path.join(WD, relp), '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
           f'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
           f'Target="{tgt}" TargetMode="External"/></Relationships>')
        rels = rels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="externalLinks/externalLink{k}.xml"/></Relationships>')
        ctx = ctx.replace("</Types>", f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>')
        order += [part, relp]
        refs.append(f'<externalReference r:id="{rid}"/>')
    wb = wb.replace("</sheets>", "</sheets><externalReferences>" + "".join(refs) + "</externalReferences>", 1)
    wr(wbp, wb); wr(rp, rels); wr(ctp, ctx)


def recalc(path):
    r = subprocess.run(["python3", "/mnt/skills/public/xlsx/scripts/recalc.py", path, "300"], capture_output=True, text=True, timeout=900)
    return r.stdout[-300:]


def values(path, sheet, refs):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    return {ref: ws[ref].value for ref in refs}


if __name__ == "__main__":
    CHK = os.path.join(WORK, "check"); shutil.rmtree(CHK, ignore_errors=True); os.makedirs(CHK)
    kids = {}
    files = sorted(f for f in os.listdir(SRC) if f.startswith("Однолинейка_") and f.endswith(".xlsx"))
    for fn in files:
        out = os.path.join(CHK, fn)
        WD, order, panel, lines = convert_child(fn, out)
        kids[fn] = dict(WD=WD, order=order, panel=panel, lines=lines)
        print("щит", fn, panel, len(lines), "линий")
    # пересчёт щитов → экспорт (строка 45 «Питание_щита»)
    EXP = ["C45", "D45", "E45", "F45", "G45", "H45", "I45"]
    for fn in files:
        recalc(os.path.join(CHK, fn))
        kids[fn]["exp"] = values(os.path.join(CHK, fn), "Питание_щита", EXP)
        print("экспорт", fn, {k: (round(v, 3) if isinstance(v, float) else v) for k, v in kids[fn]["exp"].items()})
    # ГРЩ: связи
    WDg = os.path.join(WORK, "w_grsh"); og = P7.unpack(GRSH, WDg); X.load_sst(WDg)
    ns = X.Sheet(X.sheet_path(WDg, "Нижестоящие_щиты"))
    c = {}; ext = []
    for i in range(30):
        r = 5 + i
        for col in "BC": c[f"{col}{r}"] = ce(f"{col}{r}", ns.style(f"{col}{r}"))
        for col in "DEFGHIJ": c[f"{col}{r}"] = ce(f"{col}{r}", ns.style(f"{col}{r}"))
    for i, (line, fn, _) in enumerate(LINKS):
        r = 5 + i; k = i + 1
        c[f"B{r}"] = ct(f"B{r}", line, ns.style(f"B{r}")); c[f"C{r}"] = ct(f"C{r}", fn, ns.style(f"C{r}"))
        exp = kids[fn]["exp"]
        for src_ref, col in zip(EXP, "DEFGHIJ"):
            v = exp[src_ref]
            fx = f"[{k}]Питание_щита!{src_ref}"
            c[f"{col}{r}"] = (f'<c r="{col}{r}" s="{ns.style(f"{col}{r}")}"><f>{esc(fx)}</f><v>{v}</v></c>' if isinstance(v, (int, float))
                              else f'<c r="{col}{r}" s="{ns.style(f"{col}{r}")}" t="str"><f>{esc(fx)}</f><v>{esc(v or "")}</v></c>')
        ext.append((fn, "Питание_щита", exp))
    bulk(ns, c); ns.save()
    add_ext_links(WDg, og, ext)
    gout = os.path.join(CHK, "Однолинейка ГРЩ-ЖД.xlsx")
    P7.pack(WDg, og, gout)
    print(recalc(gout))
    KN = {}
    wbv = openpyxl.load_workbook(gout, data_only=True); nsv = wbv["Нижестоящие_щиты"]
    for i, (line, fn, _) in enumerate(LINKS):
        r = 5 + i
        KN[fn] = (r, {col: nsv[f"{col}{r}"].value for col in "KLMN"}, nsv[f"O{r}"].value,
                  {col: nsv[f"{col}{r}"].value for col in "DEFGHIJ"})
        print("ГРЩ", line, KN[fn][3], KN[fn][1], KN[fn][2])
    ldv = wbv["Нагрузка_щитов"]
    print("ГРЩ блок AF5:AF12", [ldv[f"AF{r}"].value for r in range(5, 13)], "V7", ldv["V7"].value)
    # щиты: импорт от ГРЩ
    for fn in files:
        d = kids[fn]; WD = d["WD"]
        if fn in KN:
            r, kn, _, _ = KN[fn]
            ps = X.Sheet(X.sheet_path(WD, "Питание_щита")); c = {}
            cache = {}
            for col_src, col_dst in zip("KLMN", "CDEF"):
                v = kn[col_src]; src_ref = f"{col_src}{r}"; cache[src_ref] = v if v not in (None,) else ""
                fx = f"[1]Нижестоящие_щиты!{src_ref}"
                st = ps.style(f"{col_dst}6")
                c[f"{col_dst}6"] = (f'<c r="{col_dst}6" s="{st}"><f>{esc(fx)}</f><v>{v}</v></c>' if isinstance(v, (int, float))
                                    else f'<c r="{col_dst}6" s="{st}" t="str"><f>{esc(fx)}</f><v></v></c>')
            c["G6"] = ct("G6", "Однолинейка ГРЩ-ЖД.xlsx", ps.style("G6"))
            c["C12"] = ct("C12", "ГРЩ-ЖД", ps.style("C12"))
            bulk(ps, c); ps.save()
            add_ext_links(WD, d["order"], [("Однолинейка ГРЩ-ЖД.xlsx", "Нижестоящие_щиты", cache)])
        P7.pack(WD, d["order"], os.path.join(OUT, fn))
    # итог для README
    P7.pack(WDg, og, os.path.join(OUT, "Однолинейка ГРЩ-ЖД.xlsx"))
    json.dump({fn: {"panel": kids[fn]["panel"], "exp": kids[fn]["exp"]} for fn in files}, open(os.path.join(WORK, "summary.json"), "w"), ensure_ascii=False, default=str)
    print("готово")
