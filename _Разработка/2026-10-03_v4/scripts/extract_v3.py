# -*- coding: utf-8 -*-
"""extract_v3.py - вытащить ВВЕДЁННЫЕ данные из книг однолинейки v3.12 в один JSON (для сборки v4).

    py extract_v3.py --folder "<папка объекта>" --out объект.json
    py extract_v3.py <книга.xlsx> [...] --out объект.json

Только чтение (openpyxl). Формулы по умолчанию (значение «по коду») не переносятся: в v4 пустая ячейка = по коду.
Пользовательские формулы (например, E = 0,06*G) переносятся значением (кэш Excel).

Структура JSON:
  settings   - настройки (Uн, К прокл., мест на лист, мин. автомат перед БП) из первой книги
  boards     - щиты: имя, файл, схема (вводы, АВР, ИБП от), питание вручную (раздел 2), шины (B73:D82)
  places     - места схемы: щит, № места, номер линии, ввод «Однолинейки» (D,E,F,I-L,N,P,Y), фаза вручную
  lines      - линии: ввод «Исходных данных» по номеру линии
  issues     - что найдено при переносе (переименования, линии без места, дубли)
"""
import argparse
import glob
import json
import os
import sys
import warnings

import openpyxl
from openpyxl.utils import column_index_from_string as CI

warnings.filterwarnings("ignore")

# «Исходные данные»: столбец -> ключ; DEF - формула по умолчанию (не переносим), иначе формула -> значение
SD = {"C": "name", "D": "line", "E": "p", "F": "u", "G": "n", "H": "ph", "I": "l_plan", "J": "l_tray", "K": "l_pvc",
      "L": "l_pnd", "N": "cable", "O": "cores", "P": "cos", "Q": "s", "R": "qf", "S": "floor",
      "AU": "dwg_pvc", "AV": "dwg_pnd", "AW": "dwg_handles", "AX": "dwg_note", "AY": "dwg_label",
      "AZ": "kc_group", "BA": "kc", "BB": "section", "BD": "cat_pz", "BH": "k_lay", "BN": "cat", "BO": "fire"}
SD_DEF = {"F", "H", "I", "K", "L", "N", "O", "P", "Q"}       # формулы шаблона «по коду / по DWG»
DEF_TOKENS = ("VLOOKUP", "COUNT(", "IF(", "НЩ_", "INDEX(", "MATCH(")
ONE = {"D": "bus", "E": "qf_txt", "F": "rcd", "I": "dev1", "J": "dev2", "K": "dev3", "L": "module", "N": "ch",
       "P": "xt", "Y": "mode"}


def is_f(v):
    return isinstance(v, str) and v.startswith("=")


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v == "" or v.startswith("#"):
            return None
    return v


def read_book(path, issues):
    wf = openpyxl.load_workbook(path)                    # формулы
    wv = openpyxl.load_workbook(path, data_only=True)    # кэш значений
    base = os.path.basename(path)
    file_board = os.path.splitext(base)[0].replace("Однолинейка_", "").replace("Однолинейка ", "").strip()
    res = {"file": base, "file_board": file_board, "lines": [], "places": [], "boards": []}

    # --- Исходные данные
    sf, sv = wf["Исходные данные"], wv["Исходные данные"]
    for r in range(3, sf.max_row + 1):
        line = clean(sv.cell(row=r, column=4).value)
        if line is None:
            continue
        rec = {"_row": r}
        for col, key in SD.items():
            c = CI(col)
            f = sf.cell(row=r, column=c).value
            v = sv.cell(row=r, column=c).value
            if is_f(f):
                # формула шаблона (по коду / по DWG) -> не переносим; арифметика пользователя (=15+8+6) -> значение
                if col in SD_DEF and any(t in f.upper() for t in DEF_TOKENS):
                    continue
                rec[key] = clean(v)
            else:
                rec[key] = clean(f)
        rec = {k: v for k, v in rec.items() if v is not None}
        rec["line"] = str(line)
        # результаты v3.12 для сверки
        rec["_ref"] = {k: clean(sv[f"{c}{r}"].value) for k, c in (("p", "AL"), ("i", "T"), ("qf", "U"), ("s", "V"), ("l", "M"),
                       ("du", "AB"), ("du_vru", "BG"), ("grp", "BC"), ("cos", "P"), ("ikz", "BI"), ("trip", "BK"))}
        res["lines"].append(rec)

    # --- Разбивка + Однолинейка (строка = место)
    rz, ov, of = wv["Разбивка_по_щитам"], wv["Однолинейка"], wf["Однолинейка"]
    nf = wv["Нагрузка_щитов"]
    phase, phase_calc = {}, {}
    for r in range(3, nf.max_row + 1):
        d, e = clean(nf.cell(row=r, column=4).value), wf["Нагрузка_щитов"].cell(row=r, column=5).value
        if d is None:
            continue
        if e not in (None, "") and not is_f(e):
            phase[str(d)] = e
        g = clean(nf.cell(row=r, column=7).value)          # фаза (итог), как считала v3.12
        if g is not None:
            phase_calc[str(d)] = g
    for r in range(3, 1001):
        board = clean(rz.cell(row=r, column=6).value)
        if board is None:
            continue
        line = clean(rz.cell(row=r, column=8).value)
        pos = clean(ov.cell(row=r, column=2).value)
        rec = {"_row": r, "board": str(board), "pos": pos, "line": str(line) if line is not None else None}
        for col, key in ONE.items():
            f = of.cell(row=r, column=CI(col)).value
            if f is not None and not is_f(f):
                rec[key] = clean(f)
        # клемма стёрта вручную (нет ни значения, ни формулы автонумерации) -> в v4 «нет» (без номера)
        if of.cell(row=r, column=CI("P")).value is None and line is not None:
            rec["xt"] = "нет"
        if rec["line"] and rec["line"] in phase:
            rec["phase"] = phase[rec["line"]]
        if rec["line"] and rec["line"] in phase_calc:
            rec["phase_calc"] = phase_calc[rec["line"]]
        rec = {k: v for k, v in rec.items() if v is not None}
        has = any(k in rec for k in ONE.values()) or "line" in rec
        if has:
            res["places"].append(rec)

    # --- Питание_щита: щиты книги (столбцы C, D, E)
    pf, pv = wf["Питание_щита"], wv["Питание_щита"]
    for i, col in enumerate("CDE"):
        name = clean(pv[f"B{6 + i}"].value)
        if name is None:
            continue
        c = CI(col)
        man = {}
        for r, key in ((12, "from"), (13, "du_up"), (14, "ikz_min_up"), (15, "ikz_max_up"), (16, "tr_s"), (17, "tr_scheme"),
                       (18, "cable"), (19, "s"), (20, "cores"), (21, "n_par"), (22, "l"), (23, "mat"), (24, "qf_in"), (25, "qf_char")):
            v = pf.cell(row=r, column=c).value
            if v is not None and not is_f(v):
                man[key] = clean(v)
        sch = {}
        for r, key in ((65, "inputs"), (66, "avr"), (67, "ups_from")):
            v = pf.cell(row=r, column=c).value
            if v is not None and not is_f(v):
                sch[key] = clean(v)
        res["boards"].append({"name": str(name), "col": col, "manual": {k: v for k, v in man.items() if v is not None},
                              "scheme": sch})
    nfv = wv["Нагрузка_щитов"]
    for b in res["boards"]:
        r0 = {"C": 4, "D": 13, "E": 22}[b["col"]]
        b["_ref"] = {k: clean(nfv[f"{c}{r0 + d}"].value) for k, c, d in (("p_inst", "AF", 1), ("p_byt", "AF", 2), ("kc_byt", "AF", 3),
                     ("p_sil", "AF", 4), ("p_post", "AF", 5), ("pr", "AF", 6), ("kc", "AF", 7), ("ir", "AF", 8), ("cos", "V", 4))}
        b["_ref"]["du_bus"] = clean(pv[f"{b['col']}33"].value)
    buses = []
    for r in range(73, 83):
        b, s, t = (clean(pf.cell(row=r, column=c).value) for c in (2, 3, 4))
        if b and t:
            buses.append({"bus": b, "board": s, "type": t})
    res["buses"] = buses
    res["kz"] = {k: clean(pf[a].value) for k, a in (("rho_cu", "C51"), ("rho_al", "C52"), ("r_contact", "C53"), ("c_min", "C54"))}

    # --- Нижестоящие_щиты: фидеры
    nh = wv["Нижестоящие_щиты"]
    res["feeders"] = []
    for r in range(5, 35):
        b, fn = clean(nh.cell(row=r, column=2).value), clean(nh.cell(row=r, column=3).value)
        if b:
            res["feeders"].append({"line": str(b), "file": fn})

    # --- справочники книги (коды, УГО, этажи, кабели) + сколько линий книги использует код
    sp = wv["Справочная"]
    rows = list(sp.iter_rows(min_row=4, max_row=72, min_col=1, max_col=56, values_only=True))
    ref = {"codes": {}, "ugo": {}, "floors": {}, "cables": {}}
    for r in rows[:48]:
        if r[15] not in (None, ""):
            ref["codes"].setdefault(str(r[15]), [r[15], *r[16:26], r[51], r[55]])   # P..Z, AZ, BD (первая строка, как VLOOKUP)
        if r[29] not in (None, ""):
            ref["floors"].setdefault(str(r[29]), list(r[29:33]))                # AD..AG
        if r[11] not in (None, ""):
            ref["cables"].setdefault(str(r[11]), list(r[11:14]))                # L..N
    for r in rows:
        if r[46] not in (None, ""):
            ref["ugo"].setdefault(str(r[46]), list(r[46:51]))                   # AU..AY
    res["ref"] = ref
    use, fuse, cuse = {}, {}, {}
    for ln in res["lines"]:
        parts = str(ln["line"]).split(".")
        use[parts[0]] = use.get(parts[0], 0) + 1
        if len(parts) > 1:
            fuse[parts[1]] = fuse.get(parts[1], 0) + 1
        if ln.get("cable"):
            cuse[str(ln["cable"])] = cuse.get(str(ln["cable"]), 0) + 1
    res["code_use"] = use
    res["use"] = {"codes": use, "ugo": use, "floors": fuse, "cables": cuse}

    # --- настройки (Справочная)
    res["settings"] = {"places_per_sheet": clean(sp["BF4"].value), "min_qf_psu": clean(sp["BF5"].value),
                       "un_kv": clean(sp["BQ29"].value), "k_lay": clean(sp["BQ30"].value), "ko_sil": clean(sp["BQ14"].value)}
    return res


def merge(books):
    issues = []
    out = {"settings": books[0]["settings"], "kz": books[0]["kz"], "boards": [], "places": [], "lines": [],
           "buses": [], "feeders": [], "issues": issues}
    names = {}
    for bk in books:
        own = sorted({p["board"] for p in bk["places"]} | {b["name"] for b in bk["boards"]})
        ren = {}
        # книга с одним щитом, чьё имя не совпадает с именем файла -> имя файла (пример: ЩСПЗ-ЖД, внутри «ЩР-1»)
        if len(books) > 1 and len(own) == 1 and own[0] != bk["file_board"]:
            ren[own[0]] = bk["file_board"]
            issues.append(f"{bk['file']}: щит «{own[0]}» переименован в «{bk['file_board']}» (по имени файла)")
        for b in own:
            nb = ren.get(b, b)
            if nb in names and names[nb] != bk["file"]:
                issues.append(f"{bk['file']}: щит «{nb}» уже есть в {names[nb]} — переименован в «{nb}/{bk['file_board']}»")
                ren[b] = nb = f"{nb}/{bk['file_board']}"
            names[nb] = bk["file"]
        rn = lambda b: ren.get(b, b)
        for b in bk["boards"]:
            b = dict(b, name=rn(b["name"]), file=bk["file"])
            out["boards"].append(b)
        placed = set()
        for p in bk["places"]:
            p = dict(p, board=rn(p["board"]), file=bk["file"])
            out["places"].append(p)
            if p.get("line"):
                placed.add(p["line"])
        first = rn(own[0]) if own else bk["file_board"]
        for ln in bk["lines"]:
            ln = dict(ln, file=bk["file"])
            if ln["line"] not in placed and not str(ln["line"]).startswith("Ввод"):
                issues.append(f"{bk['file']}: линия «{ln['line']}» есть в «Исходных данных», но не стоит на схеме — перенесена без места в щит «{first}»")
                out["places"].append({"board": first, "pos": None, "line": ln["line"], "file": bk["file"], "unplaced": True})
            out["lines"].append(ln)
        for b in bk["buses"]:
            out["buses"].append(dict(b, board=rn(b["board"]) if b.get("board") else None, file=bk["file"]))
        for f in bk["feeders"]:
            out["feeders"].append(dict(f, from_file=bk["file"]))
    # объединённые справочники: для кода берём определение книги, где код больше всего используется
    out["ref"] = {}
    for kind in ("codes", "ugo", "floors", "cables"):
        merged, conflicts = {}, []
        keys = []
        for bk in books:
            for k in bk["ref"][kind]:
                if k not in keys:
                    keys.append(k)
        for k in keys:
            cands = [(bk["use"][kind].get(k, 0), bk["ref"][kind][k], bk["file"]) for bk in books if k in bk["ref"][kind]]
            variants = {}
            for u, v, f in cands:
                variants.setdefault(json.dumps(v, ensure_ascii=False, default=str), [0, 0, v, []])
                e = variants[json.dumps(v, ensure_ascii=False, default=str)]
                e[0] += u
                e[1] += 1
                e[3].append(f)
            best = max(variants.values(), key=lambda e: (e[0], e[1]))
            merged[k] = best[2]
            used = [e for e in variants.values() if e[0] > 0]
            if len(used) > 1:
                conflicts.append(f"справочник {kind}: «{k}» по-разному определён в книгах, где используется — взят вариант из {best[3][0]}")
        out["ref"][kind] = merged
        issues.extend(conflicts)
    # дубли номеров линий внутри объекта
    seen = {}
    for ln in out["lines"]:
        k = ln["line"]
        if k in seen and seen[k] != ln["file"] and not k.startswith("Ввод"):
            issues.append(f"номер линии «{k}» есть в {seen[k]} и {ln['file']}")
        seen.setdefault(k, ln["file"])
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("books", nargs="*")
    ap.add_argument("--folder")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    files = list(a.books)
    if a.folder:
        files += sorted(glob.glob(os.path.join(a.folder, "Однолинейка*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    # ГРЩ первым (его настройки — настройки объекта)
    files.sort(key=lambda f: (0 if "ГРЩ" in os.path.basename(f) else 1, f))
    books = []
    for f in files:
        print("чтение", os.path.basename(f), flush=True)
        books.append(read_book(f, []))
    data = merge(books)
    with open(a.out, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1, default=str)
    print(f"щитов {len(data['boards'])}, мест {len(data['places'])}, линий {len(data['lines'])}, замечаний {len(data['issues'])}")
    for i in data["issues"]:
        print(" -", i)


if __name__ == "__main__":
    main()
