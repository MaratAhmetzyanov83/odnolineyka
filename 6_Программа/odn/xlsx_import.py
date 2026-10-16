# -*- coding: utf-8 -*-
"""Импорт книги «Однолинейка» v3.12 в модель (только чтение, книга не меняется).

Берутся значения, которые Excel сохранил при последнем пересчёте (openpyxl data_only).
Адреса — как в надстройке (SX_Ribbon.bas, CollectInfo) и в шапках листов v3.12.
"""
from __future__ import annotations

import os
import warnings

import openpyxl

from .model import Book, Line, Panel, Supply

warnings.filterwarnings("ignore", module="openpyxl")

S_SRC, S_SPLIT, S_ODN, S_NET = "Исходные данные", "Разбивка_по_щитам", "Однолинейка", "Сеть_щита"
NET_SUB_ROW = 94          # «Сеть_щита», раздел Б: первая строка фидеров (30 строк)


def col(letters: str) -> int:
    """'A' -> 0, 'AB' -> 27 (индекс в кортеже строки)."""
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def is_err(v) -> bool:
    return isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A", "0")) and len(v) <= 8


def txt(v) -> str:
    if v is None or is_err(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def num(v):
    if isinstance(v, bool) or v is None or is_err(v):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").replace(" ", "").strip())
    except ValueError:
        return None


def inum(v):
    x = num(v)
    return int(round(x)) if x is not None else None


def cell(row, letters):
    i = col(letters)
    return row[i] if i < len(row) else None


def read_rows(ws, max_row=1000):
    return [r for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True)]


def read_book(path: str, rel: str) -> Book:
    book = Book(file=rel, mtime=os.path.getmtime(path))
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 — книга битая / открыта на запись / не Excel
        book.error = f"не открылась: {e}"
        return book
    try:
        if S_ODN not in wb.sheetnames or S_SRC not in wb.sheetnames:
            book.error = "не книга «Однолинейка»"
            return book
        src = read_rows(wb[S_SRC])
        odn = read_rows(wb[S_ODN])
        split = read_rows(wb[S_SPLIT]) if S_SPLIT in wb.sheetnames else []
        net = read_rows(wb[S_NET], 140) if S_NET in wb.sheetnames else []
        _fill(book, src, odn, split, net)
    finally:
        wb.close()
    return book


def _net(net, r, letters):
    """Ячейка «Сеть_щита» по номеру строки Excel."""
    return cell(net[r - 1], letters) if 0 < r <= len(net) else None


def _fill(book: Book, src, odn, split, net):
    # ---- щиты книги: «Сеть_щита» B6:B8 (порядок = столбцы C/D/E в разделе 3), иначе — по «Однолинейке»
    names = [txt(_net(net, r, "B")) for r in (6, 7, 8)] if net else []
    names = [n for n in names if n]
    if not names:
        for r in odn[2:]:
            n = txt(cell(r, "A"))
            if n and n not in names:
                names.append(n)
    panels = {}
    for k, name in enumerate(names):
        p = Panel(id=f"{book.file} / {name}", name=name, file=book.file)
        c = "CDE"[k] if k < 3 else "C"
        if net:
            r = 6 + k
            p.supply = Supply(
                from_file=txt(_net(net, r, "G")), qf=txt(_net(net, r, "H")), breaker=txt(_net(net, r, "I")),
                cable=txt(_net(net, r, "J")), length=num(_net(net, r, "K")), from_panel=txt(_net(net, r, "L")),
                in_a=num(_net(net, r, "F")), du_bus=num(_net(net, r, "C")), ikz_max=num(_net(net, r, "E")))
            if not p.supply.from_file and k == 0:
                p.supply.manual = txt(_net(net, 12, "C"))
            p.p_calc = num(_net(net, 29, c))
            p.i_calc = num(_net(net, 30, c))
            p.cos = num(_net(net, 31, c))
            p.du_bus = num(_net(net, 33, c))
            p.zp_bus = num(_net(net, 34, c))
            p.ikz_min_bus = num(_net(net, 35, c))
            p.ikz_max_bus = num(_net(net, 36, c))
            p.checks = {lbl: txt(_net(net, rr, c)) for rr, lbl in
                        ((38, "Аппарат и кабель"), (39, "Отключение КЗ"), (40, "Термостойкость"), (41, "Селективность"))}
            p.checks = {a: b for a, b in p.checks.items() if b}
            p.icu = num(_net(net, 42, c))
            p.p_inst = num(_net(net, 45 + k, "C"))     # раздел 4: строка щита (тот же порядок, что B6:B8)
            p.inputs = inum(_net(net, 65, c))
            p.avr = txt(_net(net, 66, c))
            p.category = txt(_net(net, 68, "C")) if k == 0 else ""
        panels[name] = p

    # ---- щит линии: «Разбивка_по_щитам» H (линия) -> F (щит)
    line_panel = {}
    for r in split[2:]:
        ln, pn = txt(cell(r, "H")), txt(cell(r, "F"))
        if ln and pn:
            line_panel.setdefault(ln, pn)

    # ---- «Однолинейка»: строка по номеру линии; автомат группы тянется с первой строки группы (G одинаков)
    odn_by_line = {}
    last_breaker = {}
    for i, r in enumerate(odn[2:], start=3):
        pn, qf = txt(cell(r, "A")), txt(cell(r, "G"))
        brk = txt(cell(r, "H"))
        in_ = num(cell(r, "AU"))
        if qf and brk:
            last_breaker[(pn, qf)] = (brk, in_)
        ln = txt(cell(r, "R"))
        lvl = inum(cell(r, "BU")) or 0
        if pn in panels and lvl:
            if lvl >= 2:
                panels[pn].excel_errors = (panels[pn].excel_errors or 0) + 1
            else:
                panels[pn].excel_warnings = (panels[pn].excel_warnings or 0) + 1
        if not ln:
            if pn in panels and lvl:
                panels[pn].row_notes.append(dict(row=i, qf=qf, status=txt(cell(r, "AE")), level=lvl))
            continue
        odn_by_line[ln] = dict(
            row=i, panel=pn, place=inum(cell(r, "B")), sheet=inum(cell(r, "C")), bus=txt(cell(r, "D")),
            qf=qf, breaker=brk or last_breaker.get((pn, qf), ("", None))[0],
            devices=[d for d in (txt(cell(r, x)) for x in ("I", "J", "K")) if d],
            terminal=txt(cell(r, "P")), cable=txt(cell(r, "Q")), phase=txt(cell(r, "V")),
            status=txt(cell(r, "AE")), level=lvl,
            in_a=in_ if brk else last_breaker.get((pn, qf), ("", None))[1])
    for p in panels.values():
        if p.excel_errors is None:
            p.excel_errors = 0
        if p.excel_warnings is None:
            p.excel_warnings = 0

    # ---- линии: «Исходные данные», строки с номером линии (D)
    for i, r in enumerate(src[2:], start=3):
        num_ = txt(cell(r, "D"))
        if not num_:
            continue
        o = odn_by_line.get(num_, {})
        pn = o.get("panel") or line_panel.get(num_) or (names[0] if names else "")
        ln = Line(
            number=num_, panel=pn,
            name=txt(cell(r, "C")) or txt(cell(r, "AJ")) or txt(cell(r, "B")),
            p_kw=num(cell(r, "E")), u_v=num(cell(r, "F")), count=inum(cell(r, "G")), phases=inum(cell(r, "H")),
            length_plan=num(cell(r, "I")), length=num(cell(r, "M")), cable_code=txt(cell(r, "N")),
            cores=inum(cell(r, "O")), cos=num(cell(r, "P")), s_min=num(cell(r, "Q")), in_manual=num(cell(r, "R")),
            i_calc=num(cell(r, "T")), in_a=num(cell(r, "U")), s_mm2=num(cell(r, "V")), du=num(cell(r, "AB")),
            i_start=num(cell(r, "AT")), group_ks=txt(cell(r, "BC")), du_max=num(cell(r, "BF")),
            du_sum=num(cell(r, "BG")), ikz_min=num(cell(r, "BI")), ikz_ratio=num(cell(r, "BJ")),
            ikz_check=txt(cell(r, "BK")), ikz_max=num(cell(r, "BM")), category=txt(cell(r, "BP")),
            fire=txt(cell(r, "BQ")), category_check=txt(cell(r, "BV")), low_current=inum(cell(r, "CB")) == 1,
            parallel=inum(cell(r, "CC")), cable=txt(cell(r, "AO")).replace("\n", " "), row_src=i)
        if o:
            ln.place, ln.sheet, ln.bus = o["place"], o["sheet"], o["bus"]
            ln.qf, ln.breaker, ln.devices = o["qf"], o["breaker"], o["devices"]
            ln.terminal, ln.phase, ln.status, ln.level = o["terminal"], o["phase"], o["status"], o["level"]
            ln.row_odn = o["row"]
            if o["in_a"]:
                ln.in_a = o["in_a"]          # автомат группы на схеме (может быть больше, чем по току)
            if o["cable"]:
                ln.cable = o["cable"]
        if pn not in panels:
            panels[pn] = Panel(id=f"{book.file} / {pn}", name=pn, file=book.file)
        panels[pn].lines.append(ln)

    for p in panels.values():
        p.lines.sort(key=lambda x: (x.place is None, x.place or 0, x.row_src or 0))

    # ---- раздел Б: что этот щит «знает» о нижестоящих (для сверки связей)
    if net:
        for r in range(NET_SUB_ROW, NET_SUB_ROW + 30):
            line = txt(_net(net, r, "B"))
            f = txt(_net(net, r, "C"))
            if not line and not f:
                continue
            ent = dict(line=line, file=f, p_inst=num(_net(net, r, "D")), i_calc=num(_net(net, r, "J")),
                       in_a=num(_net(net, r, "N")), check=txt(_net(net, r, "O")), row=r)
            owner = panels.get(line_panel.get(line, ""), None) or (next(iter(panels.values())) if panels else None)
            if owner:
                owner.feeders_seen.append(ent)

    book.panels = list(panels.values())
