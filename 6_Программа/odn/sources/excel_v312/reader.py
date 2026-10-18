# -*- coding: utf-8 -*-
"""Источник «книга Однолинейка v3.12»: чтение (openpyxl, значения последнего сохранения) и открытие в Excel.

Книга не меняется. Все адреса — в layout.py.
"""
from __future__ import annotations

import os
import sys
import typing
import warnings

from ...core import model as M
from ..base import Source, register
from . import layout as X

warnings.filterwarnings("ignore", module="openpyxl")


# ---------------------------------------------------------------- значения ячеек
def _col(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def _is_err(v) -> bool:
    return isinstance(v, str) and v.startswith("#") and len(v) <= 8 and v.endswith(("!", "?", "A", "0"))


def txt(v) -> str:
    if v is None or _is_err(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def num(v):
    if isinstance(v, bool) or v is None or _is_err(v):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").replace(" ", "").strip())
    except ValueError:
        return None


def inum(v):
    x = num(v)
    return int(round(x)) if x is not None else None


_CONVERT = {float: num, int: inum, str: txt, bool: lambda v: inum(v) == 1}


def _hint(cls, name):
    t = typing.get_type_hints(cls, vars(M)).get(name)
    for a in typing.get_args(t) or (t,):          # Optional[float] -> float
        if a in _CONVERT:
            return a
    return str


def set_path(obj, path: str, raw):
    """«input.length» + сырое значение ячейки -> поле нужного типа."""
    *head, last = path.split(".")
    for h in head:
        obj = getattr(obj, h)
    setattr(obj, last, _CONVERT[_hint(type(obj), last)](raw))


class Rows:
    """Лист как список строк; адрес — как в Excel (буква столбца, номер строки)."""

    def __init__(self, ws, max_row):
        self.rows = [r for r in ws.iter_rows(min_row=1, max_row=max_row, values_only=True)] if ws else []

    def get(self, row: int, letters: str):
        if not 0 < row <= len(self.rows):
            return None
        r, i = self.rows[row - 1], _col(letters)
        return r[i] if i < len(r) else None

    def data(self, first=X.FIRST_ROW):
        for i in range(first, len(self.rows) + 1):
            yield i


def _low_current(doc, src):
    """Столбец «Слаботочная линия» (patch_lowcurrent.py) — по заголовку строки 2: обычно CB, но скрипт берёт
    первый свободный от CB. Нет столбца — книга без слаботочных (правка «Слаб.» недоступна)."""
    hdr = src.rows[1] if len(src.rows) > 1 else ()
    i = next((k for k, v in enumerate(hdr) if str(v or "").strip() == X.LOW_HDR), None)
    if i is None:
        return
    doc.features.append("low_current")
    doc.low_col = _letters(i + 1)
    if doc.low_col != X.SRC_FIELDS["input.low_current"]:
        for p in doc.panels:
            for ln in p.lines:
                r = int("".join(c for c in ln.refs["input"].cell if c.isdigit()))
                set_path(ln, "input.low_current", src.get(r, doc.low_col))


def _letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _kj_rows(ws) -> list:
    """Графы 1–8 кабельного журнала — значения ячеек (сверка с отметкой «КЖ в DWG»)."""
    out = []
    for r in ws.iter_rows(min_row=X.KJ_FIRST, max_row=X.LAST_ROW, min_col=X.KJ_COL, max_col=X.KJ_COL + X.KJ_N - 1,
                          values_only=True):
        mk = str(r[0] if r and r[0] is not None else "").strip()
        if not mk or mk == "0":
            continue
        out.append([v if v is None or isinstance(v, (int, float, str)) else str(v) for v in r] + [None] * (X.KJ_N - len(r)))
    return out


# ---------------------------------------------------------------- адаптер
class ExcelV312(Source):
    kind = "excel_v312"
    extensions = (".xlsx", ".xlsm")

    def can_read(self, path):
        name = os.path.basename(path)
        return super().can_read(path) and not name.startswith(("~$", "ПЗ_"))

    def read(self, path, rel):
        import openpyxl
        doc = M.Document(file=rel, kind=self.kind, mtime=os.path.getmtime(path))
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001 — битая, не Excel, недоступна
            doc.error = f"не открылась: {e}"
            return doc
        try:
            if any(s not in wb.sheetnames for s in X.REQUIRED):
                doc.error = "не книга «Однолинейка»"
                return doc
            sh = lambda n, m=X.LAST_ROW: Rows(wb[n] if n in wb.sheetnames else None, m)  # noqa: E731
            src = sh(X.SRC)
            _fill(doc, src, sh(X.ODN), sh(X.SPLIT), sh(X.NET, X.NET_ROWS))
            _low_current(doc, src)
            if X.KJ in wb.sheetnames:
                doc.cable_journal = _kj_rows(wb[X.KJ])
        finally:
            wb.close()
        return doc

    def write(self, path, cells):
        from .writer import write_cells
        return write_cells(path, cells)

    def cells_for(self, prj, changes):
        from .writer import cells_for
        return cells_for(prj, changes)

    def open_at(self, path, ref):
        """Книга — в Excel пользователя (видимый, под его управлением), лист и ячейка выделены.
        Excel уже открыт — берём его; нет — запускаем и отдаём пользователю (UserControl), чтобы он мог сохранять
        и Excel не закрылся вместе с программой. Без pywin32 — просто открыть файл."""
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore
            pythoncom.CoInitialize()                     # запрос окна идёт в своём потоке
            try:
                xl = win32com.client.GetActiveObject("Excel.Application")
            except Exception:  # noqa: BLE001 — Excel не запущен
                xl = win32com.client.Dispatch("Excel.Application")
            xl.Visible = True
            xl.UserControl = True
            wb = next((w for w in xl.Workbooks if os.path.normcase(w.FullName) == os.path.normcase(path)), None)
            if wb is None:
                wb = xl.Workbooks.Open(path)
            if wb.ReadOnly:
                return "книга открылась только для чтения: она открыта в другом Excel или у другого пользователя"
            wb.Activate()
            if ref.sheet and ref.cell:
                ws = wb.Worksheets(ref.sheet)
                ws.Activate()
                xl.Goto(ws.Range(ref.cell), True)
            try:
                xl.WindowState = -4137                   # развернуть
            except Exception:  # noqa: BLE001
                pass
            return "ok"
        except Exception as e:  # noqa: BLE001
            if sys.platform == "win32":
                os.startfile(path)  # type: ignore[attr-defined]
                return f"открыт файл, ячейку найдите сами: {ref}"
            return f"не открыт: {e}"


register(ExcelV312())


# ---------------------------------------------------------------- разбор книги
def _row(ref) -> int:
    return int("".join(ch for ch in ref.cell if ch.isdigit()) or 0)


def _ref(doc, sheet, col, row):
    return M.Ref(doc.file, sheet, f"{col}{row}")


def _fill(doc: M.Document, src: Rows, odn: Rows, split: Rows, net: Rows):
    panels = _panels(doc, odn, net)
    names = list(panels)

    line_panel, line_split = {}, {}
    for i in split.data():
        ln, pn = txt(split.get(i, X.SPLIT_LINE)), txt(split.get(i, X.SPLIT_PANEL))
        if ln and pn:
            line_panel.setdefault(ln, pn)
        if ln:
            line_split.setdefault(ln, i)

    # «Однолинейка»: строка по номеру линии; автомат группы (G одинаков) — с первой строки группы
    odn_rows, group = {}, {}
    for i in odn.data():
        pn, qf, brk = txt(odn.get(i, X.ODN_PANEL)), txt(odn.get(i, X.ODN_QF)), txt(odn.get(i, X.ODN_BREAKER))
        if qf and brk:
            group[(pn, qf)] = (brk, num(odn.get(i, X.ODN_IN)))
        lvl = inum(odn.get(i, X.ODN_LEVEL)) or 0
        p = panels.get(pn)
        if p and lvl:
            if lvl >= 2:
                p.result.errors += 1
            else:
                p.result.warnings += 1
        ln = txt(odn.get(i, X.ODN_LINE))
        if ln:
            odn_rows[ln] = (i, pn, group.get((pn, qf), ("", None)) if qf else ("", None))
        elif p and lvl:
            p.notes.append(M.Note(lvl, txt(odn.get(i, X.ODN_STATUS)), qf, _ref(doc, X.ODN, X.ODN_STATUS, i)))

    # линии: строки «Исходных данных» с номером
    for i in src.data():
        number = txt(src.get(i, X.SRC_NUMBER))
        if not number:
            continue
        o = odn_rows.get(number)
        pn = (o[1] if o else "") or line_panel.get(number) or (names[0] if names else "")
        if pn not in panels:
            panels[pn] = M.Panel(id=f"{doc.file} / {pn}", name=pn, doc=doc.file)
        ln = M.Line(number=number, panel=panels[pn].id,
                    name=next((t for t in (txt(src.get(i, c)) for c in X.SRC_NAME) if t), ""))
        ln.result.by = "excel"
        for path, c in X.SRC_FIELDS.items():
            set_path(ln, path, src.get(i, c))
        ln.scheme.cable = ln.scheme.cable.replace("\n", " ")
        ln.refs["input"] = _ref(doc, X.SRC, X.SRC_NUMBER, i)
        if number in line_split:
            ln.refs["split"] = _ref(doc, X.SPLIT, X.SPLIT_LINE, line_split[number])
        if o:
            r, _, (brk, in_a) = o
            for path, c in X.ODN_FIELDS.items():
                v = odn.get(r, c)
                if path != "scheme.cable" or txt(v):
                    set_path(ln, path, v)
            ln.scheme.qf = txt(odn.get(r, X.ODN_QF))
            ln.scheme.breaker, ln.scheme.in_a = brk, in_a
            ln.scheme.breaker_set = txt(odn.get(r, X.ODN_BREAKER_SET))
            ln.scheme.group_head = bool(ln.scheme.qf) and bool(txt(odn.get(r, X.ODN_BREAKER)))
            ln.scheme.devices = [d for d in (txt(odn.get(r, c)) for c in X.ODN_DEVICES) if d]
            ln.result.status = txt(odn.get(r, X.ODN_STATUS))
            ln.result.level = inum(odn.get(r, X.ODN_LEVEL)) or 0
            ln.refs["scheme"] = _ref(doc, X.ODN, X.ODN_LINE, r)
            ln.refs["status"] = _ref(doc, X.ODN, X.ODN_STATUS, r)
        panels[pn].lines.append(ln)

    for p in panels.values():
        p.lines.sort(key=lambda x: (x.scheme.place is None, x.scheme.place or 0, _row(x.refs["input"])))

    # раздел Б: что щит записал о нижестоящих
    if net.rows:
        for r in range(X.NET_SUB_ROW, X.NET_SUB_ROW + X.NET_SUB_COUNT):
            fs = M.FeederSeen(ref=_ref(doc, X.NET, X.NET_SUB["p_inst"], r))
            for k, c in X.NET_SUB.items():
                set_path(fs, k, net.get(r, c))
            if not fs.line and not fs.doc:
                continue
            owner = panels.get(line_panel.get(fs.line, "")) or (next(iter(panels.values()), None))
            if owner:
                owner.feeders_seen.append(fs)

    doc.panels = list(panels.values())


def _panels(doc, odn: Rows, net: Rows):
    names = [txt(net.get(r, X.NET_SUPPLY["name"])) for r in X.NET_PANEL_ROWS]
    names = [n for n in names if n]
    if not names:
        for i in odn.data():
            n = txt(odn.get(i, X.ODN_PANEL))
            if n and n not in names:
                names.append(n)
    panels = {}
    for k, name in enumerate(names):
        p = M.Panel(id=f"{doc.file} / {name}", name=name, doc=doc.file)
        p.result.by = "excel"
        if net.rows:
            r, c = X.NET_PANEL_ROWS[min(k, 2)], X.NET_PANEL_COLS[min(k, 2)]
            for f, col in X.NET_SUPPLY.items():
                if f != "name":
                    set_path(p.supply, f, net.get(r, col))
            p.refs["supply"] = _ref(doc, X.NET, X.NET_SUPPLY["from_doc"], r)
            if not p.supply.from_doc and k == 0:
                p.supply.manual = txt(net.get(*X.NET_MANUAL_FROM))
                p.refs["supply_manual"] = _ref(doc, X.NET, X.NET_MANUAL_FROM[1], X.NET_MANUAL_FROM[0])
            for row, f in X.NET_RESULT_ROWS.items():
                set_path(p.result, f, net.get(row, c))
            p.result.checks = {lbl: txt(net.get(row, c)) for row, lbl in X.NET_CHECK_ROWS.items()
                               if txt(net.get(row, c))}
            p.refs["checks"] = _ref(doc, X.NET, c, min(X.NET_CHECK_ROWS))
            p.result.p_inst = num(net.get(X.NET_LOAD_ROW + k, "C"))
            p.scheme.icu = num(net.get(X.NET_ICU_ROW, c))
            p.scheme.is_vru = "да" if txt(net.get(X.NET_VRU_ROW, c)).lower() == "да" else ""
            p.scheme.inputs = inum(net.get(X.NET_INPUTS_ROW, c))
            p.scheme.avr = txt(net.get(X.NET_AVR_ROW, c))
            if k == 0:
                p.result.category = txt(net.get(X.NET_CAT_ROW, "C"))
            # вводы 2 и 3
            if txt(net.get(X.NET_IN23_MARK[0], X.NET_IN23_MARK[1])).upper() == X.NET_IN23_MARK[2]:
                for j in range(2):
                    r = X.NET_ROWS_IN23(k, 2 + j)
                    sp = M.Supply()
                    for f, col in X.NET_IN23.items():
                        set_path(sp, f, net.get(r, col))
                    sp.input = sp.input or (2 + j)
                    if (p.scheme.inputs or 1) >= sp.input:
                        p.inputs_more.append(sp)
        panels[name] = p
    return panels
