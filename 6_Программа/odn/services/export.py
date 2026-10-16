# -*- coding: utf-8 -*-
"""Выгрузка текущего вида таблицы окна в .xlsx (столбцы, строки — как на экране, с фильтром и сортировкой)."""
from __future__ import annotations

import io

NUM_FMT = "0.###"


def to_xlsx(title: str, columns: list, rows: list) -> bytes:
    """columns: [{t, n}], rows: [[значение, ...]] — уже отфильтровано и отсортировано окном."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or "Таблица")[:31].replace("/", "-").replace("\\", "-")
    ws.append([c.get("t", "") for c in columns])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EEF2F7")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for j, col in enumerate(columns, 1):
        width = max([len(str(col.get("t", "")))] + [len(str(r[j - 1])) for r in rows[:500] if j - 1 < len(r) and r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 6), 60)
        if col.get("n"):
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                row[0].number_format = NUM_FMT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
