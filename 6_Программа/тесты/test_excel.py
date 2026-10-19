# -*- coding: utf-8 -*-
"""Чтение книг v3.12 на тестовом объекте (make_demo.py: ГРЩ -> ЩР-1, ЩР-2 в подпапке, копия в _резерв).

Объект строится один раз (≈30 с) в папку ODN_DEMO или во временную; повторно — берётся готовый.
"""
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

from odn.services.project import load_project  # noqa: E402

FOLDER = os.environ.get("ODN_DEMO") or os.path.join(tempfile.gettempdir(), "odn_demo")


class TestExcelV312(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not os.path.isfile(os.path.join(FOLDER, "Корпус 2", "ЩР-2.xlsx")):
            import make_demo
            make_demo.main(FOLDER)
        cls.prj = load_project(FOLDER, progress=lambda s: None, use_cache=False)
        cls.p = {p.name: p for p in cls.prj.panels()}

    def test_panels(self):
        self.assertEqual(sorted(self.p), ["ГРЩ", "ЩР-1", "ЩР-2"], "копия в _резерв не должна попасть")

    def test_links(self):
        self.assertEqual((self.p["ЩР-1"].parent, self.p["ЩР-1"].parent_line), (self.p["ГРЩ"].id, "Щ.1"))
        self.assertEqual((self.p["ЩР-2"].parent, self.p["ЩР-2"].parent_line), (self.p["ГРЩ"].id, "Щ.2"))

    def test_line_fields(self):
        ln = next(x for x in self.p["ЩР-1"].lines if x.number == "R.1.01.1")
        self.assertEqual((ln.scheme.place, ln.scheme.qf, ln.scheme.breaker, ln.scheme.in_a), (19, "QFD1", "2P С16А 30мА", 16.0))
        self.assertEqual((ln.input.p_kw, ln.input.length, ln.result.in_by_current, ln.result.by), (0.06, 10.0, 6.0, "excel"))
        self.assertEqual(str(ln.refs["status"]), "Однолинейка!AE21")
        # строка группы без своего автомата берёт автомат группы
        ln2 = next(x for x in self.p["ЩР-1"].lines if x.number == "R.1.02.1")
        self.assertEqual((ln2.scheme.qf, ln2.scheme.breaker), ("QFD1", "2P С16А 30мА"))

    def test_link_issues(self):
        got = sorted((i.rule, next(p.name for p in self.prj.panels() if p.id == i.panel), i.line)
                     for i in self.prj.issues if i.group == "Связи щитов")
        self.assertEqual(got, [
            ("links.feeder_no_book", "ГРЩ", "Щ.3"),
            ("links.feeder_stale", "ЩР-2", ""),
            ("links.load_stale", "ГРЩ", "Щ.1"),
            ("links.selectivity", "ГРЩ", "Щ.1"),
            ("links.selectivity", "ГРЩ", "Щ.2"),
        ])

    def test_excel_counts(self):
        errs = [i for i in self.prj.issues if i.rule == "source.row_status" and i.level == 2 and i.panel == self.p["ЩР-1"].id]
        self.assertEqual(len(errs), self.p["ЩР-1"].result.errors)
        self.assertEqual(len(errs), 6)

    def test_cache(self):
        a = load_project(FOLDER, progress=lambda s: None)
        b = load_project(FOLDER, progress=self.fail)          # всё из кэша — ничего не читается
        self.assertEqual(len(a.issues), len(b.issues))
        self.assertEqual([p.revision for p in a.panels()], [p.revision for p in b.panels()])


class TestEdit(unittest.TestCase):
    """Запись правок, журнал, резерв, отмена — на копии объекта (тестовая запись openpyxl, без Excel)."""

    @classmethod
    def setUpClass(cls):
        TestExcelV312.setUpClass()
        cls.dir = tempfile.mkdtemp(prefix="odn_edit_")
        cls.folder = os.path.join(cls.dir, "объект")
        shutil.copytree(FOLDER, cls.folder, ignore=shutil.ignore_patterns("_окно"))
        os.environ["ODN_TEST_WRITER"] = "1"

    @classmethod
    def tearDownClass(cls):
        os.environ.pop("ODN_TEST_WRITER", None)
        shutil.rmtree(cls.dir, ignore_errors=True)

    def load(self):
        return load_project(self.folder, progress=lambda s: None)

    def test_low_current(self):
        """Слаботочная: без столбца — отказ, «включить слаботочные»; столбец по заголовку (не только CB);
        да — 1, по коду — формула назад."""
        from unittest import mock
        import openpyxl
        from odn.services import edit
        folder = os.path.join(self.dir, "слаботочные")              # своя копия: журнал общей папки не трогаем
        shutil.copytree(FOLDER, folder, ignore=shutil.ignore_patterns("_окно"))
        load = lambda: load_project(folder, progress=lambda s: None)  # noqa: E731
        prj = load()
        pid = next(p.id for p in prj.panels() if p.name == "ЩР-2")
        doc = next(d for d in prj.documents if d.file.endswith("ЩР-2.xlsx"))
        path = os.path.join(folder, doc.file)
        wb = openpyxl.load_workbook(path)
        wb["Исходные данные"]["CB2"] = None                  # книга без столбца слаботочных
        wb.save(path)
        prj = load()
        ch = [{"panel": pid, "line": prj.panel(pid).lines[0].number, "path": "input.low_current_set", "new": "да"}]
        from odn.core.changes import EditError
        with self.assertRaisesRegex(EditError, "включить слаботочные"):
            edit.apply_changes(prj, ch)
        from odn.services import book_patch
        with mock.patch.dict(os.environ, {"ODN_TEST_WRITER": "1"}):
            self.assertTrue(book_patch.run_patch(prj, doc.file, "lowcurrent")["ok"])   # «включить слаботочные в книге»
        wb = openpyxl.load_workbook(path)                    # столбец не в CB (скрипт берёт первый свободный)
        ws = wb["Исходные данные"]
        ws["CB2"], ws["CF2"] = None, "Слаботочная линия (1 — да, по коду)"
        wb.save(path)
        prj = load()
        ln = prj.panel(pid).lines[0]
        cell = ln.refs["input"].cell
        row = "".join(c for c in cell if c.isdigit())
        r = edit.apply_changes(prj, ch)
        self.assertEqual((r["written"], r["errors"]), (1, []))
        self.assertEqual(openpyxl.load_workbook(path)["Исходные данные"][f"CF{row}"].value, 1)
        self.assertTrue(load().line(pid, ln.number).input.low_current)
        r = edit.apply_changes(load(), [dict(ch[0], new="по коду")])
        self.assertEqual(r["written"], 1)
        self.assertTrue(str(openpyxl.load_workbook(path)["Исходные данные"][f"CF{row}"].value).startswith("=IF(D"))
        with self.assertRaises(EditError):
            edit.apply_changes(load(), [dict(ch[0], new="может")])

    def test_rename_line(self):
        """Номер линии-фидера: книга, «Разбивка», раздел Б, G6 нижестоящего; связь жива; отмена — всё назад."""
        import openpyxl
        from odn.services import edit
        folder = os.path.join(self.dir, "переименование")
        shutil.copytree(FOLDER, folder, ignore=shutil.ignore_patterns("_окно"))
        load = lambda: load_project(folder, progress=lambda s: None)  # noqa: E731
        prj = load()
        p = {x.name: x for x in prj.panels()}
        old = p["ЩР-1"].parent_line
        g6 = p["ЩР-1"].refs["supply"].cell
        r = edit.apply_changes(prj, [{"panel": p["ГРЩ"].id, "line": old, "path": "number", "new": "Щ.1Н"}])
        self.assertEqual(r["errors"], [])
        wb = openpyxl.load_workbook(os.path.join(folder, "ГРЩ.xlsx"), read_only=True)
        self.assertIn("Щ.1Н", [c for row in wb["Исходные данные"].iter_rows(min_col=4, max_col=4, values_only=True) for c in row])
        self.assertIn("Щ.1Н", [c for row in wb["Сеть_щита"].iter_rows(min_row=94, max_row=123, min_col=2, max_col=2, values_only=True) for c in row])
        self.assertTrue(openpyxl.load_workbook(os.path.join(folder, "ЩР-1.xlsx"))["Сеть_щита"][g6].value.endswith(" / Щ.1Н"))
        prj = load()
        p = {x.name: x for x in prj.panels()}
        self.assertEqual((p["ЩР-1"].parent, p["ЩР-1"].parent_line), (p["ГРЩ"].id, "Щ.1Н"))
        with self.assertRaisesRegex(Exception, "уже есть"):
            edit.apply_changes(prj, [{"panel": p["ГРЩ"].id, "line": "Щ.1Н", "path": "number", "new": "Щ.2"}])
        prj = load()
        self.assertEqual(edit.undo_last(prj)["errors"], [])
        prj = load()
        p = {x.name: x for x in prj.panels()}
        self.assertEqual(p["ЩР-1"].parent_line, old)

    def test_write_undo(self):
        from odn.services import edit
        prj = self.load()
        pid = next(p.id for p in prj.panels() if p.name == "ЩР-1")
        r = edit.apply_changes(prj, [
            {"panel": pid, "line": "R.1.01.1", "path": "input.length_plan", "new": "12,5"},
            {"panel": pid, "line": "R.1.01.1", "path": "scheme.breaker_set", "new": "1P C10А"},
            {"panel": pid, "line": "R.1.02.1", "path": "input.cores", "new": 5},
        ])
        self.assertEqual((r["written"], r["docs"], r["errors"]), (3, ["ЩР-1.xlsx"], []))
        prj = self.load()
        self.assertEqual(prj.line(pid, "R.1.01.1").input.length_plan, 12.5)
        self.assertEqual(prj.line(pid, "R.1.01.1").scheme.breaker_set, "1P C10А")
        self.assertEqual(prj.line(pid, "R.1.02.1").input.cores, 5)
        j = [r for r in edit.read_journal(prj) if r.get("kind") == "правка"]
        self.assertEqual(len(j), 1)
        self.assertEqual(j[0]["changes"][0]["cell"], "Исходные данные!I21")
        self.assertTrue(os.path.isfile(os.path.join(self.folder, j[0]["backup"])))
        # ничего не изменилось — ничего не пишется
        self.assertEqual(edit.apply_changes(prj, [{"panel": pid, "line": "R.1.01.1", "path": "input.length_plan", "new": 12.5}])["written"], 0)
        # отмена
        r = edit.undo_last(prj)
        self.assertEqual(r["written"], 3)
        prj = self.load()
        self.assertEqual((prj.line(pid, "R.1.01.1").input.length_plan, prj.line(pid, "R.1.01.1").scheme.breaker_set), (8.0, "авто"))
        self.assertEqual(edit.undo_last(prj)["errors"], ["нечего отменять"])
        # служебная папка не читается как книга объекта
        self.assertEqual(sorted(p.name for p in prj.panels()), ["ГРЩ", "ЩР-1", "ЩР-2"])

    def test_links_program_and_books(self):
        from odn.core.changes import EditError
        from odn.services import link_writer, structure
        prj = self.load()
        P = {p.name: p.id for p in prj.panels()}
        # вариант А: связь в программе — ЩР-2 от ГРЩ по Щ.3 (в книге — Щ.2)
        structure.operate(prj, {"op": "set_link", "panel": P["ЩР-2"], "parent": P["ГРЩ"], "line": "Щ.3"})
        prj = self.load()
        p2 = prj.panel(P["ЩР-2"])
        self.assertEqual((p2.parent_line, p2.link_src, p2.book_line), ("Щ.3", "программа", "Щ.2"))
        self.assertEqual(prj.line(P["ГРЩ"], "Щ.3").feeds, P["ЩР-2"])
        self.assertIn(("links.mismatch", P["ЩР-2"]), {(i.rule, i.panel) for i in prj.issues})
        with self.assertRaises(EditError):                                   # цикл: ГРЩ от ЩР-1
            structure.operate(prj, {"op": "set_link", "panel": P["ГРЩ"], "parent": P["ЩР-1"], "line": ""})
        # вариант Б: та же связь — в книги; связь программы снимается
        r = link_writer.write_link(prj, P["ЩР-2"], P["ГРЩ"], "Щ.3")
        self.assertTrue(r["ok"])
        structure.operate(prj, {"op": "clear_link", "panel": P["ЩР-2"]})
        prj = self.load()
        p2 = prj.panel(P["ЩР-2"])
        self.assertEqual((p2.parent_line, p2.link_src, p2.book_line), ("Щ.3", "книга", "Щ.3"))
        self.assertNotIn(("links.mismatch", P["ЩР-2"]), {(i.rule, i.panel) for i in prj.issues})
        # отмена записи связи — книги из копий
        self.assertTrue(link_writer.undo_last(prj)["ok"])
        prj = self.load()
        self.assertEqual(prj.panel(P["ЩР-2"]).parent_line, "Щ.2")
        self.assertFalse(link_writer.undo_last(prj)["ok"])

    def test_conflict(self):
        from odn.services import edit
        prj = self.load()
        pid = next(p.id for p in prj.panels() if p.name == "ЩР-2")
        path = os.path.join(self.folder, prj.panel(pid).doc)
        os.utime(path, (os.path.getatime(path), os.path.getmtime(path) + 10))     # книгу изменили после чтения
        r = edit.apply_changes(prj, [{"panel": pid, "line": "R.1.01.1", "path": "input.p_kw", "new": 1}])
        self.assertEqual(r["written"], 0)
        self.assertIn("изменилась после чтения", r["errors"][0])

    def test_locked(self):
        from unittest import mock
        from odn.services import edit
        from odn.sources.excel_v312 import writer
        prj = self.load()
        pid = next(p.id for p in prj.panels() if p.name == "ГРЩ")
        book = os.path.join(self.folder, "ГРЩ.xlsx")
        lock = os.path.join(self.folder, "~$ГРЩ.xlsx")
        name = "Иванов"
        with open(lock, "wb") as f:                     # метка Excel: длина + ANSI, с байта 55 — длина + UTF-16
            f.write(bytes([len(name)]) + name.encode("cp1251").ljust(54, b" ") + bytes([len(name), 0]) + name.encode("utf-16-le"))
        try:
            self.assertEqual(writer.lock_owner(lock), name)
            self.assertEqual(writer.lock_reason(book), "")       # старая метка после сбоя Excel — не мешает
            with mock.patch.object(writer, "_writable", return_value=False):
                r = edit.apply_changes(prj, [{"panel": pid, "line": "Щ.1", "path": "input.p_kw", "new": 9}])
                self.assertIn("открыта в Excel у пользователя «Иванов»", r["errors"][0])
                os.remove(lock)
                self.assertIn("Excel без окна", writer.lock_reason(book))
        finally:
            if os.path.exists(lock):
                os.remove(lock)

    def test_export(self):
        from odn.services import export
        import io
        import openpyxl
        data = export.to_xlsx("Линии", [{"t": "Линия"}, {"t": "P", "n": True}], [["R.1", 0.5], ["R.2", None]])
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual([[c.value for c in r] for r in ws.iter_rows()], [["Линия", "P"], ["R.1", 0.5], ["R.2", None]])




class TestExclude(unittest.TestCase):
    """Резервные книги: слова в имени, свой список исключений, предупреждение о двойном щите."""

    def test_exclude(self):
        TestExcelV312.setUpClass()
        from odn.services.project import add_exclusion, skipped
        d = tempfile.mkdtemp(prefix="odn_excl_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt"))
            shutil.copy(os.path.join(obj, "ЩР-1.xlsx"), os.path.join(obj, "ЩР-1 резерв.xlsx"))     # пропустится по слову
            shutil.copy(os.path.join(obj, "ЩР-1.xlsx"), os.path.join(obj, "ЩР-1 старая.xlsx"))     # не пропустится
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            docs = sorted(x.file for x in prj.documents)
            self.assertNotIn("ЩР-1 резерв.xlsx", docs)
            self.assertIn("ЩР-1 старая.xlsx", docs)
            dup = [i for i in prj.issues if i.rule == "links.duplicate_panel"]
            self.assertEqual(len(dup), 2)
            add_exclusion(obj, "ЩР-1 старая.xlsx")
            prj = load_project(obj, progress=lambda s: None)
            self.assertNotIn("ЩР-1 старая.xlsx", [x.file for x in prj.documents])
            self.assertFalse([i for i in prj.issues if i.rule == "links.duplicate_panel"])
            self.assertTrue(skipped("Корпус 2/ЩР-2.xlsx", ["корпус 2/*"]))
            self.assertTrue(skipped("a/b.xlsx", ["b.xlsx"]))
            self.assertFalse(skipped("a/b.xlsx", ["c.xlsx"]))
        finally:
            shutil.rmtree(d, ignore_errors=True)




class TestExternal(unittest.TestCase):
    """Связи в другие папки: открыта только «Корпус 2» — ГРЩ (выше) и ЩР-1 (его нижестоящий) читаются по связям."""

    def test_external(self):
        TestExcelV312.setUpClass()
        d = tempfile.mkdtemp(prefix="odn_ext_")
        try:
            obj = os.path.join(d, "комплекс")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "_резерв"))
            sub = os.path.join(obj, "Корпус 2")
            prj = load_project(sub, progress=lambda s: None, use_cache=False)
            docs = {x.file.replace("\\", "/"): x.external for x in prj.documents}
            self.assertEqual(docs, {"ЩР-2.xlsx": False, "../ГРЩ.xlsx": True, "../ЩР-1.xlsx": True})
            p = {x.name: x for x in prj.panels()}
            self.assertEqual((p["ЩР-2"].parent, p["ЩР-1"].parent), (p["ГРЩ"].id, p["ГРЩ"].id))
            # запись в книгу из другой папки: резервная копия — внутри _окно/резерв
            from odn.services import edit
            os.environ["ODN_TEST_WRITER"] = "1"
            try:
                r = edit.apply_changes(prj, [{"panel": p["ГРЩ"].id, "line": "Щ.1", "path": "input.length_plan", "new": 33}])
            finally:
                os.environ.pop("ODN_TEST_WRITER", None)
            self.assertEqual(r["written"], 1, r)
            j = [x for x in edit.read_journal(prj) if x.get("kind") == "правка"][-1]
            self.assertTrue(j["backup"].startswith(os.path.join("_окно", "резерв")) and ".." not in j["backup"], j["backup"])
        finally:
            shutil.rmtree(d, ignore_errors=True)




class TestInputs(unittest.TestCase):
    """Вводы 2 и 3: ЩР-2 — ввод 2 от ГРЩ по линии Щ.3 («Сеть_щита» C65 = 2, AA6)."""

    def test_second_input(self):
        TestExcelV312.setUpClass()
        import openpyxl
        d = tempfile.mkdtemp(prefix="odn_in2_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt"))
            path = os.path.join(obj, "Корпус 2", "ЩР-2.xlsx")
            wb = openpyxl.load_workbook(path)
            net = wb["Сеть_щита"]
            net["C65"], net["U4"], net["V6"], net["AA6"], net["AM6"], net["AO6"] = 2, "ВВОДЫ 2 И 3", 2, "..\\ГРЩ.xlsx / Щ.3", "3P C16А", 20
            wb.save(path)
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            p = {x.name: x for x in prj.panels()}
            sp = p["ЩР-2"].inputs_more
            self.assertEqual([(x.input, x.parent, x.parent_line, x.breaker, x.length) for x in sp],
                             [(2, p["ГРЩ"].id, "Щ.3", "3P C16А", 20.0)])
            ln = prj.line(p["ГРЩ"].id, "Щ.3")
            self.assertEqual((ln.feeds, ln.feeds_input), (p["ЩР-2"].id, 2))
            self.assertNotIn(("links.feeder_no_book", "Щ.3"), {(i.rule, i.line) for i in prj.issues})
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestDrawingStamp(unittest.TestCase):
    """Отметки SXDRAW (<книга>.drawn.txt) -> «однолинейка устарела»."""

    def test_stamps(self):
        TestExcelV312.setUpClass()
        import openpyxl
        from odn.services import drawing_stamp as ds
        d = tempfile.mkdtemp(prefix="odn_drawn_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "*.drawn.txt"))
            book = os.path.join(obj, "ГРЩ.xlsx")
            load = lambda: {p.name: p for p in load_project(obj, progress=lambda s: None, use_cache=False).panels()}  # noqa: E731
            p = load()
            codes = lambda x: "(" + " ".join(str(ord(c)) for c in x) + ")"  # noqa: E731
            when = ds.book_time(os.path.getmtime(book))
            with open(os.path.join(obj, "ГРЩ.drawn.txt"), "w", encoding="ascii") as f:
                f.write(";; odn drawn v1\n")
                f.write(f'(DRAWN {codes("ГРЩ")} {codes("C:\\Чертежи\\ЭМ.dwg")} "{when}" "-" "2026-10-18 10:00:00" T)\n')
                # cp1251 (старый AutoLISP): тот же щит, другой чертёж, время книги — неизвестная версия
                f.write(f'(DRAWN ({" ".join(str(b) for b in "ГРЩ".encode("cp1251"))}) (65 46 100 119 103) "2001-01-01 00:00:00" "-" '
                        f'"2026-10-18 09:00:00" T)\n')
            p = load()
            st = {x["dwg"]: x["state"] for x in p["ГРЩ"].drawings}
            self.assertEqual(st, {"C:\\Чертежи\\ЭМ.dwg": "ok", "A.dwg": "check"})
            # книгу пересохранили без изменений — схема актуальна
            os.utime(book, (os.path.getmtime(book) + 120,) * 2)
            self.assertEqual(load()["ГРЩ"].drawings[0]["state"], "ok")
            # данные схемы изменились — устарела, замечание
            wb = openpyxl.load_workbook(book)
            ws = wb["Исходные данные"]
            r = next(r for r in range(3, 300) if ws[f"D{r}"].value and ws[f"E{r}"].value)
            ws[f"E{r}"] = ws[f"E{r}"].value + 1          # Pу линии
            wb.save(book)
            os.utime(book, (os.path.getmtime(book) + 240,) * 2)
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            g = {x.name: x for x in prj.panels()}["ГРЩ"]
            self.assertEqual(g.drawings[0]["state"], "stale")
            self.assertTrue(any(i.rule == "drawings.stale" and i.panel == g.id for i in prj.issues))
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestKJStamp(unittest.TestCase):
    """Отметка «КЖ в DWG» (<книга>.kj.txt) -> «КЖ устарел» с перечнем изменённых кабелей."""

    def test_kj(self):
        TestExcelV312.setUpClass()
        import openpyxl
        from odn.services.drawing_stamp import same_cell
        self.assertTrue(same_cell("30", 30.456) and same_cell("0,5%", 0.005) and same_cell("#Н/Д", "#N/A"))
        self.assertFalse(same_cell("31", 30.4))
        d = tempfile.mkdtemp(prefix="odn_kj_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "*.drawn.txt"))
            book = os.path.join(obj, "ГРЩ.xlsx")
            wb = openpyxl.load_workbook(book)
            ws = wb.create_sheet("Кабельный_журнал")
            rows = [("W1", "ГРЩ", "ЩР-1", "ВВГнг", "5х16", 30.456, "", ""), ("W2", "ГРЩ", "ЩР-2", "ВВГнг", "5х10", 45, "", "")]
            for i, r in enumerate(rows):
                for k, v in enumerate(r):
                    ws.cell(row=5 + i, column=16 + k, value=v)
            wb.save(book)
            with open(os.path.join(obj, "ГРЩ.kj.txt"), "w", encoding="utf-8-sig") as f:
                f.write("DWG\tD:\\КЖ_ГРЩ.dwg\nAT\t2026-10-18 21:00:00\nSAVED\t1\n")
                f.write("ROW\tW1\tГРЩ\tЩР-1\tВВГнг\t5х16\t30\t\t\nROW\tW2\tГРЩ\tЩР-2\tВВГнг\t5х10\t45\t\t\n")
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            doc = next(x for x in prj.documents if x.file == "ГРЩ.xlsx")
            self.assertEqual((doc.kj["state"], doc.kj["cables"]), ("ok", 2))
            ws = wb["Кабельный_журнал"]
            ws["T6"] = "5х16"                                   # W2: сечение (графа 5)
            ws.cell(row=7, column=16, value="W3")               # новый кабель
            wb.save(book)
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            doc = next(x for x in prj.documents if x.file == "ГРЩ.xlsx")
            self.assertEqual(doc.kj["state"], "stale")
            self.assertIn("W2 (гр. 5: 5х10 → 5х16)", doc.kj["why"])
            self.assertIn("новые 1: W3", doc.kj["why"])
            self.assertTrue(any(i.rule == "drawings.kj_stale" for i in prj.issues))
            # обновить КЖ из окна (тест: без Excel и AutoCAD — только отметка) -> «актуален», DWG — из отметки
            from odn.services import kj_writer
            os.environ["ODN_TEST_WRITER"] = "1"
            try:
                res = kj_writer.issue(prj, [{"doc": "ГРЩ.xlsx", "dwg": "КЖ_ГРЩ"}, {"doc": "нет.xlsx", "dwg": ""}])
            finally:
                os.environ.pop("ODN_TEST_WRITER", None)
            self.assertEqual([x["dwg"] for x in res["done"]], [os.path.abspath(os.path.join(obj, "КЖ_ГРЩ.dwg"))])
            self.assertEqual(res["errors"], ["нет.xlsx: нет книги"])
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            doc = next(x for x in prj.documents if x.file == "ГРЩ.xlsx")
            self.assertEqual((doc.kj["state"], doc.kj["cables"]), ("ok", 3))
            self.assertEqual(kj_writer._target(prj, doc, ""), doc.kj["dwg"])          # «обновить» — DWG из отметки
            self.assertEqual(kj_writer.default_dwg(doc), "КЖ_ГРЩ.dwg")
            # убрать кабель из КЖ (<книга>.kj.skip.txt): КЖ устарел («убраны»), после выпуска — актуален без него
            self.assertEqual(kj_writer.set_skip(prj, "ГРЩ.xlsx", [" W3 ", ""])["n"], 1)
            self.assertEqual(kj_writer.cables(prj, "ГРЩ.xlsx")["skip"], ["W3"])
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            doc = next(x for x in prj.documents if x.file == "ГРЩ.xlsx")
            self.assertEqual(doc.kj["state"], "stale")
            self.assertIn("убраны 1: W3", doc.kj["why"])
            os.environ["ODN_TEST_WRITER"] = "1"
            try:
                kj_writer.issue(prj, [{"doc": "ГРЩ.xlsx", "dwg": ""}])
            finally:
                os.environ.pop("ODN_TEST_WRITER", None)
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            doc = next(x for x in prj.documents if x.file == "ГРЩ.xlsx")
            self.assertEqual((doc.kj["state"], doc.kj["cables"]), ("ok", 2))
            kj_writer.set_skip(prj, "ГРЩ.xlsx", [])
            self.assertFalse(os.path.exists(os.path.join(obj, "ГРЩ.kj.skip.txt")))
            self.assertEqual(kj_writer._target(prj, doc, "сх\\КЖ"), os.path.abspath(os.path.join(obj, "сх", "КЖ.dwg")))
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestDismiss(unittest.TestCase):
    """Замечание снято вручную (_окно/снятые.json): уходит из списка и счёта, «вернуть» — снова в списке."""

    def test_dismiss(self):
        TestExcelV312.setUpClass()
        from odn.core.changes import EditError
        from odn.services import dismiss
        from odn.services.project import recheck
        d = tempfile.mkdtemp(prefix="odn_dis_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "*.drawn.txt"))
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            i = next(x for x in prj.issues if x.line)
            k, n = dismiss.key(i), len(prj.issues)
            res = dismiss.dismiss(prj, [{"rule": i.rule, "panel": i.panel, "line": i.line, "text": i.text}], "проверено")
            self.assertEqual(res["n"], 1)
            recheck(prj)
            self.assertNotIn(k, {dismiss.key(x) for x in prj.issues})
            self.assertEqual(len(prj.issues), n - 1)
            self.assertEqual([dismiss.key(x) for x in prj.dismissed], [k])
            self.assertTrue(prj.dismissed[0].dismissed.endswith(": проверено"))
            prj = load_project(obj, progress=lambda s: None, use_cache=False)       # после перечитывания — тоже снято
            self.assertEqual(len(prj.dismissed), 1)
            with self.assertRaises(EditError):                                      # такого замечания уже нет
                dismiss.dismiss(prj, [{"rule": i.rule, "panel": i.panel, "line": i.line, "text": "другой текст"}])
            dismiss.restore(prj, [{"rule": i.rule, "panel": i.panel, "line": i.line, "text": i.text}])
            recheck(prj)
            self.assertEqual((len(prj.issues), prj.dismissed), (n, []))
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestFixes(unittest.TestCase):
    """«Исправить ошибки»: предложения правок по замечаниям."""

    def test_propose(self):
        TestExcelV312.setUpClass()
        from odn.services import fixes, structure
        from odn.services.project import recheck
        self.assertEqual(fixes._no_rcd("2P C16А 30мА"), "1P C16А")
        self.assertEqual(fixes.pick_in(25.6), 32)
        d = tempfile.mkdtemp(prefix="odn_fix_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "*.drawn.txt"))
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            p = {x.name: x for x in prj.panels()}
            got = {(x["kind"], x["panel_name"], x["line"]): x for x in fixes.propose(prj)}
            # замечания самой книги: «5 жил на 1 фазе» -> 3 жилы; у замечаний — пояснение
            cores = [x for x in got.values() if x["rule"] == "source.row_status" and x["path"] == "input.cores"]
            self.assertTrue(cores and all(x["now"] == "3" and x["checked"] for x in cores))
            self.assertTrue(all(i.hint for i in prj.issues if "жил на 1 фазе" in i.text))
            from odn.checks.hints import hint_for
            self.assertIn("Iдоп", hint_for("КАБЕЛЬ НЕ ЗАЩИЩЁН"))
            fx = got[("edit", "ГРЩ", p["ЩР-1"].parent_line)]
            self.assertEqual((fx["path"], fx["now"]), ("input.in_manual", "32"))   # 1,6 × 16 А -> 32 А
            # УЗО над линиями ЩР-1 — их дифавтоматы лишние: предложение «без 30 мА», не отмечено сразу
            heads = [ln.number for ln in p["ЩР-1"].lines if "мА" in ln.scheme.breaker and (ln.scheme.group_head or not ln.scheme.qf)][:2]
            structure.operate(prj, {"op": "add_node", "panel": p["ЩР-1"].id, "kind": "УЗО", "lines": heads})
            recheck(prj)
            rcd = [x for x in fixes.propose(prj) if x["rule"] == "nodes.extra_rcd"]
            self.assertEqual(sorted(x["line"] for x in rcd), sorted(heads))
            self.assertTrue(all(x["path"] == "scheme.breaker_set" and "мА" not in x["now"] and not x["checked"] for x in rcd))
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestPhasesFeeders(unittest.TestCase):
    """Фазы с нижестоящими щитами; кабель фидера в двух книгах (лоток / ПВХ / ПНД)."""

    def test_phases_and_cable(self):
        TestExcelV312.setUpClass()
        import openpyxl
        d = tempfile.mkdtemp(prefix="odn_ph_")
        try:
            obj = os.path.join(d, "объект")
            shutil.copytree(FOLDER, obj, ignore=shutil.ignore_patterns("_окно", "*.nodes.txt", "*.drawn.txt"))
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            p = {x.name: x for x in prj.panels()}
            g = p["ГРЩ"]
            for k in ("L1", "L2", "L3"):          # ГРЩ = свои линии + ЩР-1 + ЩР-2
                self.assertAlmostEqual(g.phase_i[k], g.phase_own[k] + p["ЩР-1"].phase_i[k] + p["ЩР-2"].phase_i[k], places=2)
            feeder = prj.line(g.id, p["ЩР-1"].parent_line)
            self.assertEqual(feeder.phase_i, p["ЩР-1"].phase_i)
            # строка ввода в ЩР-1: гофра ПНД 9 м, а у фидера в ГРЩ — другое
            path = os.path.join(obj, "ЩР-1.xlsx")
            wb = openpyxl.load_workbook(path)
            ws = wb["Исходные данные"]
            r = next(r for r in range(3, 21) if not ws[f"D{r}"].value)
            ws[f"D{r}"], ws[f"K{r}"], ws[f"L{r}"], ws[f"M{r}"] = "Ввод", 0, (feeder.input.pnd or 0) + 9, feeder.input.length
            wb.save(path)
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            p = {x.name: x for x in prj.panels()}
            feeder = prj.line(p["ГРЩ"].id, p["ЩР-1"].parent_line)
            self.assertEqual(list(feeder.mismatch), ["input.pnd"])
            self.assertEqual(list(prj.line(p["ЩР-1"].id, "Ввод").mismatch), ["input.pnd"])
            self.assertTrue(any(i.rule == "links.cable_mismatch" and "гофра ПНД" in i.text for i in prj.issues))
            # «Исправить ошибки»: строка ввода = фидер; запись в книгу — и расхождения нет
            from unittest import mock
            from odn.services import edit, fixes
            fx = [x for x in fixes.propose(prj) if x["rule"] == "links.cable_mismatch"]
            self.assertEqual([(x["line"], x["path"]) for x in fx], [("Ввод", "input.pnd")])
            with mock.patch.dict(os.environ, {"ODN_TEST_WRITER": "1"}):
                r = edit.apply_changes(prj, [{k: fx[0][k] for k in ("panel", "line", "path")} | {"new": fx[0]["now"]}])
            self.assertEqual((r["written"], r["errors"]), (1, []))
            prj = load_project(obj, progress=lambda s: None, use_cache=False)
            self.assertFalse(any(i.rule == "links.cable_mismatch" for i in prj.issues))
        finally:
            shutil.rmtree(d, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
