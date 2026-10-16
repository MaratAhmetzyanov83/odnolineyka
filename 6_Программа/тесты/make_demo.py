# -*- coding: utf-8 -*-
"""Тестовый объект из книги-примера (только значения, без формул): ГРЩ -> ЩР-1, ЩР-2 (в подпапке).

    python make_demo.py <папка>

Намеренные ошибки для проверки программы:
  - ГРЩ считает нагрузку ЩР-1 по старой (Руст 3 кВт);
  - ЩР-2 знает старый In фидера (16 А, а у ГРЩ 20 А) и фидер неселективен (20 < 1,6 × 16);
  - у ГРЩ фидер «Питание ЩР-3», книги которого нет;
  - копия в «_резерв» не должна попасть в объект.
"""
import os
import shutil
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
EXAMPLE = os.path.join(HERE, "..", "..", "1_Таблицы_Excel", "Однолинейка_пример_ЖК-Остров.xlsx")
SHEETS = ("Исходные данные", "Разбивка_по_щитам", "Однолинейка", "Сеть_щита")


def base():
    wb = openpyxl.load_workbook(EXAMPLE, data_only=True)
    for n in list(wb.sheetnames):
        if n not in SHEETS:
            del wb[n]
    return wb


def rename(wb, name):
    for r in range(3, 1001):
        if wb["Однолинейка"].cell(r, 1).value == "ЩР":
            wb["Однолинейка"].cell(r, 1).value = name
        if wb["Разбивка_по_щитам"].cell(r, 6).value == "ЩР":
            wb["Разбивка_по_щитам"].cell(r, 6).value = name
    wb["Сеть_щита"]["B6"] = name


def feeder(wb, row, number, name, in_a, length):
    """Линия «Исходных данных» и «Однолинейки» в строке row -> фидер щита."""
    src, odn = wb["Исходные данные"], wb["Однолинейка"]
    src[f"D{row}"] = number
    src[f"C{row}"] = name
    src[f"AJ{row}"] = name
    src[f"BC{row}"] = "Щит"
    src[f"U{row}"] = in_a
    src[f"M{row}"] = length
    odn[f"R{row}"] = number
    odn[f"S{row}"] = name
    odn[f"H{row}"] = f"3P C{in_a}А"
    odn[f"AU{row}"] = in_a
    odn[f"G{row}"] = f"QF{row}"
    odn[f"AE{row}"] = None
    odn[f"BU{row}"] = 0
    for r in range(3, 1001):
        if wb["Разбивка_по_щитам"].cell(r, 8).value == src.cell(row, 1).value:
            wb["Разбивка_по_щитам"].cell(r, 8).value = number


def main(folder):
    os.makedirs(os.path.join(folder, "Корпус 2"), exist_ok=True)
    os.makedirs(os.path.join(folder, "_резерв"), exist_ok=True)
    # ГРЩ
    wb = base()
    rename(wb, "ГРЩ")
    net = wb["Сеть_щита"]
    net["C12"] = "ТП"
    feeder(wb, 21, "Щ.1", "Питание ЩР-1", 25, 30)
    feeder(wb, 22, "Щ.2", "Питание ЩР-2", 20, 45)
    feeder(wb, 23, "Щ.3", "Питание ЩР-3", 16, 20)
    net["B94"], net["C94"], net["D94"], net["N94"] = "Щ.1", "ЩР-1.xlsx", 3.0, 25
    net["B95"], net["C95"], net["D95"], net["N95"] = "Щ.2", "Корпус 2\\ЩР-2.xlsx", 5.52, 20
    p = os.path.join(folder, "ГРЩ.xlsx")
    wb.save(p)
    shutil.copy(p, os.path.join(folder, "_резерв", "ГРЩ.xlsx"))
    # ЩР-1
    wb = base()
    rename(wb, "ЩР-1")
    net = wb["Сеть_щита"]
    net["G6"], net["F6"], net["K6"], net["L6"], net["H6"] = "ГРЩ.xlsx / ГРЩ", 25, 30, "ГРЩ", "QF21"
    wb.save(os.path.join(folder, "ЩР-1.xlsx"))
    # ЩР-2
    wb = base()
    rename(wb, "ЩР-2")
    net = wb["Сеть_щита"]
    net["G6"], net["F6"], net["K6"], net["L6"], net["H6"] = "..\\ГРЩ.xlsx / ГРЩ", 16, 45, "ГРЩ", "QF22"
    wb.save(os.path.join(folder, "Корпус 2", "ЩР-2.xlsx"))
    print("готово:", folder)


if __name__ == "__main__":
    main(sys.argv[1])
