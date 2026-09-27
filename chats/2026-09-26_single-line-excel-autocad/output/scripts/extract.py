import pickle, re, collections, json
import openpyxl
SP = "/tmp/claude-0/-home-claude/5e645175-0e52-5350-b586-c0cb5927e030/scratchpad/"
rows = pickle.load(open(SP + "ins.pkl", "rb"))
segs, texts = pickle.load(open(SP + "geom.pkl", "rb"))
band = [(p, t.strip(), h) for p, t, h in texts if 1470 < p[1] < 1580 and t.strip()]
ins = [r for r in rows if 1300 < r[3] < 1700]
vsegs = [(p, q) for p, q in segs if abs(p[0] - q[0]) < 0.05 and max(p[1], q[1]) > 1535]

wv = openpyxl.load_workbook("/root/.claude/uploads/5e645175-0e52-5350-b586-c0cb5927e030/aa5313f2-__________________.xlsx", data_only=True)
rz = wv["Разбивка_по_щитам"]

out = []
for i in range(0, 391):
    xc = 55 + 20 * i + 10
    x0, x1 = xc - 10, xc + 10
    T = [(p[0] - xc, p[1], t) for p, t, h in band if x0 <= p[0] < x1]
    B = [r for r in ins if x0 <= r[2] < x1]
    d = {"i": i, "row": 3 + i, "line": rz.cell(3 + i, 8).value, "name": rz.cell(3 + i, 10).value,
         "cable": rz.cell(3 + i, 16).value}
    qf = [r for r in B if r[0] == "QF1"]
    if qf:
        a = qf[0][6]
        lab = [t for dx, y, t in T if t in ("QF", "QFD", "QSD") and y > 1520]
        num = [t for dx, y, t in T if re.fullmatch(r"\d+", t) and 1525 < y < 1552 and dx < 0]
        mA = any(t == a.get("МА") for dx, y, t in T)
        # bus: top of vertical at xc
        tops = [max(p[1], q[1]) for p, q in vsegs if abs(p[0] - xc) < 0.8]
        d["qf"] = {"P": a.get("P"), "A": a.get("A"), "mA": a.get("МА") if mA else "", "lab": lab[:1], "num": num[:1], "top": round(max(tops), 1) if tops else None}
    devs = []
    for r in B:
        if r[0] in ("MK-5-1", "Кмка", "Реверс", "Реле напряжения схема", "Блок питания_BERKER", "Драйвер", "км-3"):
            devs.append(r[0])
    if any(t == "T" for dx, y, t in T if y > 1500):
        devs.append("T")
    d["dev"] = devs
    mods = [t for dx, y, t in T if re.fullmatch(r"(R|Di|DA|DI|R)\.?\d+", t) and 1490 < y < 1500]
    model = [t for dx, y, t in T if re.match(r"Z[A-Z0-9\-]+", t) and 1488 < y < 1500]
    for r in B:
        if r[0] in ("24-1", "Диммер 4 канала", "далишка", "24-4"):
            model += [v for v in r[6].values() if v]
    d["mod"] = mods[:1]; d["model"] = model[:1]
    ch = [t for dx, y, t in T if re.fullmatch(r"\d+", t) and 1496.5 < y < 1500 and dx < 0]
    d["ch"] = ch[:1]
    has_xt = any(t == "XT" for dx, y, t in T if 1476 < y < 1485)
    xtn = [t for dx, y, t in T if re.fullmatch(r"\d+", t) and 1478.5 < y < 1485 and dx > 1]
    d["xt"] = xtn[:1] if has_xt else []
    out.append(d)
json.dump(out, open(SP + "slots.json", "w"), ensure_ascii=False, indent=0)
for d in out:
    s = f"{d['i']:3} r{d['row']:3} {str(d['line'] or ''):18.18}"
    if "qf" in d: s += f" QF[{''.join(d['qf']['lab'])}{''.join(d['qf']['num'])} {d['qf']['P']} {d['qf']['A']} {d['qf']['mA']} top={d['qf']['top']}]"
    if d["dev"]: s += f" DEV{d['dev']}"
    if d["mod"] or d["model"]: s += f" MOD{d['mod']}{d['model']}"
    if d["ch"]: s += f" ch{d['ch'][0]}"
    if d["xt"]: s += f" XT{d['xt'][0]}"
    print(s)
