# Строит отчёт «Длины трасс ЖК остров.xlsx» из result.json (первый прогон по 03_Чертежи.dwg).
# result.json формирует прототип (extract.py + match.py): группы -> [(уровень, хэндл, слой, тип линии, длина м, как найдено)],
# multi (полилинии с несколькими группами), orphans (полилинии без группы), issues (проблемы выносок).
import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open("result.json"))
LT = {"JIS_02_0.7": "JIS_02_0.7", "Continuous": "Continuous", "ШТРИХОВАЯ": "ШТРИХОВАЯ", "Lotok": "Lotok"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Длины по группам"
F = "Arial"; hdr = Font(name=F, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="305496")
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
heads = ["Номер линии", "Уровень", "Кол-во полилиний", "JIS_02_0.7, м", "Continuous, м", "ШТРИХОВАЯ, м", "Lotok, м",
         "Итого по плану, м", "Как найдено", "Хэндлы полилиний"]
ws.append(heads)

def key(g):
    return [int(x) if x.isdigit() else x for x in re.split(r'[.\-]', g)]

r = 2
for g in sorted(d["groups"], key=key):
    segs = d["groups"][g]
    lv = sorted({s[0] for s in segs})
    sums = {k: 0.0 for k in LT}
    for s in segs: sums[s[3]] = sums.get(s[3], 0) + s[4]
    how = "выноска" if all(s[5] == "выноска" for s in segs) else "выноска + цепочка"
    ws.append([g, ", ".join(lv), len(segs)] + [round(sums[k], 2) or None for k in LT] +
              [f"=SUM(D{r}:G{r})", how, ", ".join(s[1] for s in segs)])
    r += 1
for c in range(1, len(heads) + 1):
    ws.cell(1, c).font = hdr; ws.cell(1, c).fill = fill
    ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
for row in ws.iter_rows(min_row=1, max_row=r - 1):
    for c in row:
        c.border = B
        if c.row > 1: c.font = Font(name=F)
        if c.column in (4, 5, 6, 7, 8) and c.row > 1: c.number_format = "0.00"
for i, w in enumerate([16, 9, 11, 13, 13, 13, 10, 15, 18, 40], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32; ws.freeze_panes = "A2"
n = r
ws.cell(n + 1, 1, "Примечания").font = Font(name=F, bold=True)
notes = ["Длины — только по плану (горизонталь), в метрах, из модели по окнам основных видовых экранов (Ур.1 и Ур.2).",
         "Спуски, подъём на +4.320 и запас сюда не входят — они считаются формулой в «Исходных данных».",
         "Какой тип линии = пол / потолок / лоток — нужно подтвердить; колонки названы по типу линии как в DWG.",
         "Хэндл можно найти в AutoCAD: (sssetfirst nil (ssadd (handent \"830A6E\")))"]
for i, t in enumerate(notes): ws.cell(n + 2 + i, 1, t).font = Font(name=F, italic=True, color="595959")

ws2 = wb.create_sheet("Проверить")
ws2.append(["Уровень", "Что не так", "Номер линии / хэндл", "Подробности"])
for lvl, h, layer, lt, L, gs in d["multi"]:
    ws2.append([lvl, "На одну полилинию указывают выноски разных групп", h,
                f"{', '.join(gs)}; {layer}; {lt}; {L} м — длина не отнесена ни к одной"])
for lvl, h, layer, lt, L in d["orphans"]:
    ws2.append([lvl, "Полилиния без группы (нет выноски и не продолжает другую трассу)", h, f"{layer}; {lt}; {L} м"])
for i in d["issues"]:
    ws2.append([i[0], i[1], i[2], f"выноска {i[3]}"])
for c in range(1, 5):
    ws2.cell(1, c).font = hdr; ws2.cell(1, c).fill = fill
for row in ws2.iter_rows(min_row=2):
    for c in row: c.font = Font(name=F); c.border = B
for i, w in enumerate([9, 58, 20, 70], 1): ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
wb.save("Длины трасс ЖК остров.xlsx")
