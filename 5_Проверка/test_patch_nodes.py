# -*- coding: utf-8 -*-
"""Проверка формул patch_nodes.py без Excel: маленькая книга с теми же столбцами «Однолинейки», пересчёт LibreOffice.

    python test_patch_nodes.py            — нужен soffice (LibreOffice) и openpyxl; код возврата 0 — всё сошлось

Формулы-«хозяева» (Z, AR, AS, AJ, H, AU, AE, G) — упрощённые копии формул v3.12 той же структуры; к ним применяются
те же функции fix_* из patch_nodes.py, новые столбцы CL…CS — как в книге. Целую книгу LibreOffice 24 не считает
(LET, SINGLE), поэтому так. Итог в Excel проверяется на копии книги.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import warnings

import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import patch_nodes as P  # noqa: E402

warnings.filterwarnings("ignore")
XLFN = ("MAXIFS", "MINIFS", "TEXTJOIN")


def r1c1_to_a1(f: str, row: int, col: int) -> str:
    def ref(m):
        rr, cc = m.group(1), m.group(2)
        r = row if rr is None else (row + int(rr[1:-1]) if rr.startswith("[") else int(rr))
        c = col if cc is None else (col + int(cc[1:-1]) if cc.startswith("[") else int(cc))
        rd = "" if rr is None or rr.startswith("[") else "$"
        cd = "" if cc is None or cc.startswith("[") else "$"
        return f"{cd}{cl(c)}{rd}{r}"
    parts = re.split(r'("[^"]*")', f)                     # строки в кавычках не трогаем
    for i in range(0, len(parts), 2):
        parts[i] = re.sub(r"(?<![A-Za-z_])R(\[-?\d+\]|\d+)?C(\[-?\d+\]|\d+)?(?![A-Za-z_(])", ref, parts[i])
        for fn in XLFN:
            parts[i] = re.sub(rf"(?<![\w.]){fn}\(", f"_xlfn.{fn}(", parts[i])
    return "".join(parts)


def a1_xlfn(f):
    parts = re.split(r'("[^"]*")', f)
    for i in range(0, len(parts), 2):
        for fn in XLFN:
            parts[i] = re.sub(rf"(?<![\w.]){fn}\(", f"_xlfn.{fn}(", parts[i])
    return "".join(parts)


# упрощённые формулы v3.12 (R1C1); номера столбцов — как в книге
HOST = {
    "AF": '=IF(RC5<>"",COUNTIFS(R3C1:RC1,RC1,R3C5:RC5,"?*"),IF(RC1<>R[-1]C1,"",R[-1]C))',
    "AG": '=IF(RC36<>"",RC36,IF(RC1<>R[-1]C1,"",R[-1]C))',
    "Z": '=IF(RC5="","",SUMIFS(R3C21:R1000C21,R3C32:R1000C32,RC32,R3C1:R1000C1,RC1,R3C40:R1000C40,0))',
    "AA": '=IF(RC5="","",IFERROR(MINIFS(Ser!R1C1:R30C1,Ser!R1C1:R30C1,">="&N(RC26)),""))',
    "AO": '=IF(RC5="","",N(RC27))',
    "AQ": '=IF(RC5="","",RC41)',
    "AR": '=IF(RC5="","",MAX(1,MAXIFS(R3C37:R1000C37,R3C32:R1000C32,RC32,R3C1:R1000C1,RC1)))',
    "AS": '=IF(RC5="","",IF(AND(RC5<>"авто",RC5<>"+"),IF(ISNUMBER(SEARCH("мА",RC5)),1,0),IF(RC6="да",1,'
          'IF(RC6="нет",0,IF(MAXIFS(R3C39:R1000C39,R3C32:R1000C32,RC32,R3C1:R1000C1,RC1)>0,1,0)))))',
    "AT": '=IF(RC5="","",IF(RC44>=3,IF(RC45=1,"4P","3P"),IF(RC45=1,"2P","1P")))',
    "H": '=IF(RC5="","",IF(AND(RC5<>"авто",RC5<>"+"),RC5,RC46&" С"&RC43&"А"&IF(RC45=1," 30мА","")))',
    "AU": '=IF(RC8="","",IFERROR(--MID(SUBSTITUTE(RC8,"С","C"),FIND("C",SUBSTITUTE(RC8,"С","C"))+1,'
          'FIND("А",RC8,FIND("C",SUBSTITUTE(RC8,"С","C")))-FIND("C",SUBSTITUTE(RC8,"С","C"))-1),""))',
    "AJ": '=IF(RC5="","",IF(RC45=1,"QFD","QF")&RC32)',
    "G": '=IF(LEFT(RC8,3)="QS ",SUBSTITUTE(IF(RC18<>"",RC33,""),"QF","QS"),IF(RC18<>"",RC33,""))',
    "AE": '=IF(AND(RC18="",RC7=""),"",TEXTJOIN("; ",TRUE,IF(AND(RC38="R",RC5<>"",N(RC45)=0),"розетка без 30 мА","")))',
}
# строки: место, R линия, D, E, F, U ток, V фаза, AK фаз, AM 30мА-код, AL код, CE (узел КМ), CL (тип узла)
ROWS = [
    (1, "R.1", "", "авто", "", 5, "L1", 1, 1, "R", "", ""),       # 3: обычная розетка -> QFD1
    (2, "", "", "авто", "", None, "", None, 0, "", "", "УЗО"),     # 4: узел У1 (УЗО)
    (3, "R.2", "У1", "авто", "", 7, "L1", 1, 1, "R", "", ""),     # 5: розетка под УЗО -> QF, без замечания
    (4, "R.3", "У1", "авто", "", 3, "L2", 1, 1, "R", "", ""),     # 6
    (5, "", "У1", "авто", "", None, "", None, 0, "", "", "Автомат"),  # 7: узел У2 под У1
    (6, "L.1", "У2", "авто", "", 2, "L3", 1, 0, "L", "", ""),     # 8
    (7, "R.4", "", "", "", 1.5, "L3", 1, 1, "R", "", ""),         # 9: та же группа, что L.1 (E пусто)
    (8, "P.1", "", "авто", "", 10, "L1,L2,L3", 3, 0, "P", "", ""),  # 10: обычная 3ф
    (9, "", "", "авто", "", None, "", None, 0, "", "KM-1", ""),   # 11: узел КМ
    (10, "V.1", "KM-1", "авто", "", 4, "L2", 1, 0, "V", "", ""),  # 12
    (11, "U.1", "ИБП", "авто", "", 2, "L1", 1, 0, "U", "", ""),   # 13: вторая шина ИБП
    (12, "", "", "QS 3P 63А", "", None, "", None, 0, "", "", "Рубильник"),  # 14: узел У3 без линий
]


def build(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = P.ODN
    ser = wb.create_sheet("Ser")
    for i, v in enumerate([6, 10, 16, 20, 25, 32, 40, 50, 63, 80, 100, 125], 1):
        ser.cell(i, 1, v)
    ws["A2"], ws["AF2"], ws["CO2"], ws["CQ2"] = "Щит", "гр.", "путь", "30 мА выше"
    n = len(ROWS)
    for i, (pl, r, d, e, f, u, v, ak, am, al, ce, typ) in enumerate(ROWS):
        row = 3 + i
        vals = {"A": "ЩР", "B": pl, "R": r or None, "D": d or None, "E": e or None, "F": f or None, "U": u,
                "V": v or None, "AK": ak, "AM": am, "AL": al or None, "AN": 0, "CE": ce or None, "CL": typ or None}
        for c, x in vals.items():
            ws[f"{c}{row}"] = x
        for c, f1 in HOST.items():
            fn = P.ODN_FIX.get(c)
            g = fn(f1) if fn else f1
            assert g, c
            ws[f"{c}{row}"] = r1c1_to_a1(g, row, ci(c))
        for c, (_, fa) in P.COLS.items():
            if fa:
                ws[f"{c}{row}"] = a1_row(fa, row)
    wb.save(path)
    return n


def a1_row(fa, row):
    """Формула строки 3 -> строка row (как заполнение вниз в Excel): относительные номера строк сдвигаются."""
    def sh(m):
        col, dollar, r = m.group(1), m.group(2), int(m.group(3))
        return f"{col}{dollar}{r if dollar else r + row - 3}"
    parts = re.split(r'("[^"]*")', fa)
    for i in range(0, len(parts), 2):
        parts[i] = re.sub(r"(\$?[A-Z]{1,3})(\$?)(\d+)\b", sh, parts[i])
    return a1_xlfn("".join(parts))


def main():
    d = tempfile.mkdtemp(prefix="nodes_")
    src = os.path.join(d, "t.xlsx")
    build(src)
    out = os.path.join(d, "out")
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, src],
                   check=True, capture_output=True, timeout=180)
    res = openpyxl.load_workbook(os.path.join(out, "t.xlsx"), data_only=True)[P.ODN]
    g = lambda c, r: res[f"{c}{r}"].value  # noqa: E731
    want = [
        # (столбец, строка, ожидание)
        ("CM", 4, "У1"), ("CM", 7, "У2"), ("CM", 14, "У3"), ("CM", 3, None),
        ("CO", 3, "/"), ("CO", 5, "/У1/"), ("CO", 7, "/У1/"), ("CO", 8, "/У1/У2/"), ("CO", 9, "/У1/У2/"),
        ("CO", 12, "/KM-1/"), ("CO", 13, "/ИБП/"),
        ("CN", 4, 7), ("CN", 7, 3.5),
        ("Z", 4, 7), ("Z", 7, 3.5), ("AQ", 4, 10), ("AQ", 7, 6),
        ("AR", 4, 1), ("AS", 3, 1), ("AS", 4, 1), ("AS", 5, 0), ("AS", 6, 0), ("AS", 7, 0), ("AS", 8, 0),
        ("CQ", 5, 1), ("CQ", 8, 1), ("CQ", 9, 1), ("CQ", 3, 0), ("CQ", 10, 0),
        ("AJ", 3, "QFD1"), ("AJ", 4, "QD2"), ("AJ", 5, "QF3"), ("AJ", 7, "QF5"),
        ("H", 4, "УЗО 2P 10А 30мА"), ("H", 7, "1P С6А"), ("H", 5, "1P С10А"), ("H", 14, "QS 3P 63А"),
        ("AU", 4, 10), ("AU", 7, 6), ("AU", 14, 63),
        ("G", 4, "QD2"), ("G", 7, "QF5"), ("G", 14, "QS11"),
        ("CS", 5, None), ("CS", 8, None), ("CS", 12, None), ("CS", 13, "ИБП"), ("CS", 3, None),
    ]
    fails = []
    # повторный запуск: доработанная формула узнаётся по метке и не меняется
    for c, fn in P.ODN_FIX.items():
        fixed = fn(HOST[c])
        if (P.MARK if c != "AE" else "RC96") not in fixed:
            fails.append(f"{c}: после доработки нет метки — повторный запуск испортит формулу")
    if P.fix_by("=IF(N(RC76)=0,\"\",INDEX(Однолинейка!R3C4:R1000C4,RC76))").count("R3C97:R1000C97") != 1:
        fails.append("BY: D -> CS не заменён")
    for c, r, w in want:
        v = g(c, r)
        if isinstance(w, (int, float)) and isinstance(v, (int, float)):
            ok = abs(v - w) < 1e-9
        else:
            ok = (v in (None, "") and w is None) or v == w
        if not ok:
            fails.append(f"{c}{r}: {v!r}, нужно {w!r}")
    cr4, cr7, cr14 = g("CR", 4), g("CR", 7), g("CR", 14)
    if "сумма автоматов под ним" not in str(cr4):
        fails.append(f"CR4 (УЗО 10 А, автоматы 10+6+6): {cr4!r}")
    if "неселективно" not in str(cr7):
        fails.append(f"CR7 (автомат 6 А над 6 А): {cr7!r}")
    if "нет линий" not in str(cr14):
        fails.append(f"CR14 (рубильник без линий): {cr14!r}")
    ae = {r: str(g("AE", r) or "") for r in range(3, 15)}
    if "розетка без 30 мА" in ae[5] or "розетка без 30 мА" in ae[6]:
        fails.append(f"AE5/AE6: под УЗО не нужно 30 мА — {ae[5]!r} / {ae[6]!r}")
    if "сумма автоматов" not in ae[4]:
        fails.append(f"AE4 — нет проверки узла: {ae[4]!r}")
    shutil.rmtree(d, ignore_errors=True)
    print("ОК" if not fails else "ОШИБКИ:\n  " + "\n  ".join(fails))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
